"""P0-10 Excel 核验报告：把 ValidationResult 导出为 .xlsx。

本模块只消费 P0-08 的统一结果模型，不读取业务 Excel，
不重新实现任何核验逻辑。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from excel_qc.models import ValidationResult

_SUMMARY_HEADER = ["项目", "内容"]
_DETAIL_HEADER = [
    "序号",
    "文件",
    "Sheet",
    "行号",
    "列号",
    "列字母",
    "字段名",
    "单元格",
    "错误类型",
    "错误编码",
    "实际值",
    "期望值",
    "错误原因",
    "来源",
]
_SUMMARY_COLUMN_WIDTHS = (18, 40)
_DETAIL_COLUMN_WIDTHS = (7, 26, 16, 8, 8, 10, 16, 12, 16, 14, 18, 18, 44, 16)

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


def export_validation_report(
    result: ValidationResult,
    output_path: str | Path,
) -> Path:
    """把 ValidationResult 导出为 Excel 核验报告，返回输出文件路径。

    - 只消费 ValidationResult，不打开、不修改被核验的源 Excel；
    - issues 按 ValidationResult 中的既定顺序写入；
    - 输出路径不允许与源文件相同，避免覆盖用户原始数据；
    - 输出文件已存在时由调用方指定的输出路径覆盖。
    """
    output = Path(output_path)
    if output.resolve() == result.file.resolve():
        raise ValueError(f"核验报告不能覆盖源文件: {output}")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "核验汇总"
    _write_summary(summary_sheet, result)

    detail_sheet = workbook.create_sheet("错误明细")
    _write_detail(detail_sheet, result)

    workbook.save(output)
    return output


def _write_summary(sheet: Worksheet, result: ValidationResult) -> None:
    summary = result.summary
    rows = [
        ("核验文件", result.file.name),
        (
            "结论",
            "全部通过，无错误。"
            if summary.error_count == 0
            else f"发现 {summary.error_count} 条错误，需人工处理。",
        ),
        ("总检查数", summary.total_checks),
        ("通过数", summary.pass_count),
        ("错误数", summary.error_count),
        ("格式错误", summary.format_error_count),
        ("数据不存在", summary.data_not_found_count),
        ("数据错误", summary.data_error_count),
        ("跳过空行", summary.skipped_empty_row_count),
    ]
    _append_header(sheet, _SUMMARY_HEADER)
    for row_index, (label, value) in enumerate(rows, start=2):
        _write_row(sheet, row_index, [label, value])
        sheet.cell(row_index, 1).font = Font(bold=True)
        sheet.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
    _set_column_widths(sheet, _SUMMARY_COLUMN_WIDTHS)


def _write_detail(sheet: Worksheet, result: ValidationResult) -> None:
    _append_header(sheet, _DETAIL_HEADER)
    for row_index, issue in enumerate(result.issues, start=2):
        location = issue.location
        _write_row(
            sheet,
            row_index,
            [
                row_index - 1,
                location.source_file.name,
                location.sheet_name,
                location.row_number,
                location.column_number,
                location.column_letter,
                location.column_name,
                location.cell_reference,
                issue.error_type.value,
                issue.error_code,
                issue.actual_value,
                issue.expected_value,
                issue.reason,
                issue.source.value,
            ],
        )
    _set_column_widths(sheet, _DETAIL_COLUMN_WIDTHS)
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def _append_header(sheet: Worksheet, header: list[str]) -> None:
    _write_row(sheet, 1, header)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _BORDER
    sheet.freeze_panes = "A2"


def _write_row(sheet: Worksheet, row_index: int, values: list[object]) -> None:
    for column_index, value in enumerate(values, start=1):
        cell = sheet.cell(row_index, column_index, value)
        cell.border = _BORDER


def _set_column_widths(sheet: Worksheet, widths: list[int | float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
