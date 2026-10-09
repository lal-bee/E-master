"""P1-10 回源核验报告：只消费 P1-09 内存结果。"""

from __future__ import annotations

import json
import os
import tempfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel_qc.source_verification import (
    FieldComparisonStatus,
    NormalizedFieldValue,
    SourceVerificationResult,
    VerificationIssue,
    VerificationLocation,
    VerificationRecord,
    VerificationRowStatus,
)

MAX_CELL_TEXT = 32767
MAX_SHEET_ROWS = 1048576


class SourceVerificationReportError(Exception):
    """回源核验报告写入异常。"""


class SourceVerificationReportValueError(SourceVerificationReportError):
    """结果值无法无损写入 Excel。"""


class SourceVerificationReportCapacityError(SourceVerificationReportError):
    """报告超出 Excel 容量。"""


_STATUS_LABELS = {
    VerificationRowStatus.MATCH: "一致",
    VerificationRowStatus.MISSING_IN_SYSTEM: "系统缺失",
    VerificationRowStatus.EXTRA_IN_SYSTEM: "系统新增",
    VerificationRowStatus.FIELD_CHANGED: "字段变化",
    VerificationRowStatus.DUPLICATE_KEY: "重复主键",
    VerificationRowStatus.UNVERIFIABLE: "无法核验",
}

_SUMMARY_HEADER = ("项目", "内容")
_STATS_HEADER = ("指标", "状态码", "数量", "单位", "口径说明")
_DETAIL_HEADER = (
    "主键字段", "主键值（含类型）", "结果类型", "结果码", "明细层级",
    "标准字段", "字段结果码", "源原始值", "源原始值类型", "源标准化值",
    "源标准化值类型", "源字段状态", "源应用规则", "系统原始值",
    "系统原始值类型", "系统标准化值", "系统标准化值类型", "系统字段状态",
    "系统应用规则", "比较规则", "原因", "源文件", "源Sheet", "源行号",
    "源列号", "源列字母", "源单元格", "系统文件", "系统Sheet", "系统行号",
    "系统列号", "系统列字母", "系统单元格",
)
_ISSUE_HEADER = (
    "问题类别", "问题编码", "问题所属侧", "主键值（含类型）", "标准字段",
    "原因", "阻断核验", "文件", "Sheet", "物理行号", "物理列号",
    "列字母", "单元格", "原字段名", "来源侧", "关联位置序号",
)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_STATUS_FILLS = {
    "FIELD_CHANGED": "FFF2CC",
    "MISSING_IN_SYSTEM": "FCE4D6",
    "EXTRA_IN_SYSTEM": "FCE4D6",
    "DUPLICATE_KEY": "F8D7DA",
    "UNVERIFIABLE": "F8D7DA",
}


def export_source_verification_report(
    result: SourceVerificationResult,
    output_path: str | Path,
    *,
    overwrite: bool = False,
) -> Path:
    """导出 P1-09 结果为四 Sheet .xlsx；默认拒绝覆盖，失败不替换已有文件。"""
    if not isinstance(result, SourceVerificationResult):
        raise TypeError("result 必须是 SourceVerificationResult")
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        raise SourceVerificationReportError("报告输出路径必须是 .xlsx 文件")
    for input_path in (result.original_source.source_file, result.system_export.source_file):
        if _same_file(output, input_path):
            raise SourceVerificationReportError(f"报告不能覆盖输入文件：{input_path}")
    if output.exists() and not overwrite:
        raise FileExistsError(f"报告已存在，未启用覆盖：{output}")

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    workbook: Workbook | None = None
    try:
        workbook = _build_workbook(result)
        fd, name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=".xlsx", dir=output.parent)
        os.close(fd)
        temporary = Path(name)
        workbook.save(temporary)
        if overwrite:
            os.replace(temporary, output)
        else:
            # 同目录硬链接使用独占创建语义，避免检查之后被并发创建的目标遭覆盖。
            os.link(temporary, output)
            temporary.unlink()
        temporary = None
        return output
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _same_file(first: Path, second: Path) -> bool:
    if os.path.normcase(str(first.resolve(strict=False))) == os.path.normcase(str(second.resolve(strict=False))):
        return True
    try:
        return os.path.samefile(first, second)
    except (FileNotFoundError, OSError):
        return False


def _build_workbook(result: SourceVerificationResult) -> Workbook:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "核验概要"
    statistics = workbook.create_sheet("差异统计")
    details = workbook.create_sheet("差异明细")
    issues = workbook.create_sheet("异常记录")
    _write_overview(overview, result)
    _write_statistics(statistics, result)
    _write_details(details, result)
    _write_issues(issues, result)
    return workbook


def _write_overview(sheet: Worksheet, result: SourceVerificationResult) -> None:
    summary = result.summary
    if not summary.verification_complete:
        conclusion = "无法判定"
    elif summary.data_consistent:
        conclusion = "一致"
    else:
        conclusion = "不一致"
    rows = (
        ("原始源文件", str(result.original_source.source_file)),
        ("系统导出文件", str(result.system_export.source_file)),
        ("报告生成时间", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("核验时间", "P1-09 结果未提供"),
        ("核验完成状态", "已完成" if summary.verification_complete else "未完成"),
        ("核验完成状态代码", str(summary.verification_complete)),
        ("数据一致性结论", conclusion),
        ("data_consistent", str(summary.data_consistent)),
        ("原始源输入行数", summary.original_source_input_row_count),
        ("系统导出输入行数", summary.system_export_input_row_count),
        ("唯一键匹配对数", summary.matched_row_count),
        ("一致结果记录数", summary.unchanged_row_count),
        ("变化结果记录数", summary.changed_row_count),
        ("系统缺失记录数", summary.missing_in_system_row_count),
        ("系统新增记录数", summary.extra_in_system_row_count),
        ("重复键冲突组数", summary.duplicate_key_group_count),
        ("无法核验结果记录数", summary.unverifiable_row_count),
        ("变化字段数", summary.changed_field_count),
        ("问题数量", summary.issue_count),
        ("异常提示", "存在阻断或无法判定问题，需查看异常记录" if not summary.verification_complete else "无阻断问题"),
        ("已知限制", "仅对配置字段按 P1-09 规则核验；真实业务样本验证尚未完成；不推断 KEY_CHANGED"),
    )
    _header(sheet, _SUMMARY_HEADER)
    for row in rows:
        _append(sheet, row)
    _widths(sheet, (25, 82))


def _write_statistics(sheet: Worksheet, result: SourceVerificationResult) -> None:
    summary = result.summary
    rows = (
        ("原始源输入", "ORIGINAL_SOURCE_INPUT", summary.original_source_input_row_count, "输入行", "原始源文档数据行"),
        ("系统导出输入", "SYSTEM_EXPORT_INPUT", summary.system_export_input_row_count, "输入行", "系统导出数据行"),
        ("唯一键匹配", "MATCHED_PAIRS", summary.matched_row_count, "匹配对", "包含匹配后字段不可判定的对"),
        ("一致", "MATCH", summary.unchanged_row_count, "结果记录", "配置字段全部一致"),
        ("系统缺失", "MISSING_IN_SYSTEM", summary.missing_in_system_row_count, "结果记录", "源侧有效唯一键在系统侧不存在"),
        ("系统新增", "EXTRA_IN_SYSTEM", summary.extra_in_system_row_count, "结果记录", "系统侧有效唯一键在源侧不存在"),
        ("字段变化", "FIELD_CHANGED", summary.changed_row_count, "结果记录", "一行多个字段变化只计一次"),
        ("重复主键", "DUPLICATE_KEY", summary.duplicate_key_group_count, "冲突组", "一组可含双方多条输入行"),
        ("无法核验", "UNVERIFIABLE", summary.unverifiable_row_count, "结果记录", "可与匹配对数交叉，不与其他单位相加"),
        ("已比较字段", "COMPARED_FIELDS", summary.compared_field_count, "字段项", "仅含一致和变化字段"),
        ("一致字段", "EQUAL_FIELDS", summary.equal_field_count, "字段项", "逐项比较一致"),
        ("变化字段", "CHANGED_FIELDS", summary.changed_field_count, "字段项", "包括不可判定行中已确认变化的字段"),
        ("无法比较字段", "UNVERIFIABLE_FIELDS", summary.unverifiable_field_count, "字段项", "字段不存在或处理/类型失败等"),
        ("问题总数", "ISSUES", summary.issue_count, "问题对象", "一项问题可关联多个来源位置"),
    )
    _header(sheet, _STATS_HEADER)
    for row in rows:
        _append(sheet, row)
    for code, count in summary.issue_counts:
        _append(sheet, ("问题类型", code, count, "问题对象", "按 P1-09 问题编码统计"))
    _widths(sheet, (20, 28, 12, 15, 58))


def _write_details(sheet: Worksheet, result: SourceVerificationResult) -> None:
    _header(sheet, _DETAIL_HEADER)
    for record in result.records:
        if record.status is VerificationRowStatus.MATCH:
            continue
        visible = tuple(
            difference for difference in record.field_differences
            if difference.status is not FieldComparisonStatus.EQUAL
        )
        if visible:
            for difference in visible:
                _append(sheet, _detail_row(record, difference))
        else:
            _append(sheet, _detail_row(record, None))
    _widths(sheet, (24, 42, 16, 23, 12, 20, 20, 25, 15, 25, 15, 18, 25,
                    25, 15, 25, 15, 18, 25, 18, 44, 34, 18, 12, 12, 12, 14,
                    34, 18, 12, 12, 12, 14))
    _filter(sheet)


def _detail_row(record: VerificationRecord, difference: Any) -> tuple[Any, ...]:
    source = difference.original_source if difference else None
    system = difference.system_export if difference else None
    source_location = source.location if source else _first(record.original_source_rows)
    system_location = system.location if system else _first(record.system_export_rows)
    source_values = _field_cells(source)
    system_values = _field_cells(system)
    return (
        json.dumps(record.key_fields, ensure_ascii=False), _key_text(record.key_values),
        _STATUS_LABELS[record.status], record.status.value,
        "字段" if difference else "行", difference.standard_field_name if difference else None,
        difference.status.value if difference else None,
        *source_values, *system_values,
        difference.comparison_rule.value if difference else None,
        difference.reason if difference else _STATUS_LABELS[record.status],
        *_location_cells(source_location), *_location_cells(system_location),
    )


def _field_cells(field: NormalizedFieldValue | None) -> tuple[str | None, ...]:
    if field is None:
        return (None,) * 6
    return (
        _value_text(field.original_value), _type_name(field.original_value),
        _value_text(field.normalized_value), _type_name(field.normalized_value),
        field.status.value, json.dumps(field.applied_rules, ensure_ascii=False),
    )


def _write_issues(sheet: Worksheet, result: SourceVerificationResult) -> None:
    _header(sheet, _ISSUE_HEADER)
    for issue in result.issues:
        locations: Iterable[VerificationLocation | None] = issue.locations or (None,)
        for index, location in enumerate(locations, start=1):
            _append(sheet, _issue_row(issue, location, index if location else None))
    _widths(sheet, (17, 30, 19, 42, 20, 66, 14, 34, 18, 13, 13, 12, 15, 20, 18, 16))
    _filter(sheet)


def _issue_row(issue: VerificationIssue, location: VerificationLocation | None, index: int | None) -> tuple[Any, ...]:
    return (
        issue.category.value, issue.code, issue.side.value, _key_text(issue.key_values),
        issue.standard_field_name, issue.message, "是" if issue.blocks_verification else "否",
        location.source_file if location else None,
        location.sheet_name if location else None,
        location.row_number if location else None,
        location.column_number if location else None,
        location.column_letter if location else None,
        location.cell_address if location else None,
        location.source_field_name if location else None,
        location.side.value if location else issue.side.value,
        index,
    )


def _first(locations: tuple[VerificationLocation, ...]) -> VerificationLocation | None:
    return locations[0] if locations else None


def _location_cells(location: VerificationLocation | None) -> tuple[Any, ...]:
    if location is None:
        return (None,) * 6
    return (
        str(location.source_file), location.sheet_name, location.row_number,
        location.column_number, location.column_letter, location.cell_address,
    )


def _key_text(values: tuple[Any, ...]) -> str:
    return json.dumps(
        [{"type": _type_name(value), "value": _value_text(value)} for value in values],
        ensure_ascii=False,
    )


def _value_text(value: Any) -> str:
    if value is None:
        return "None"
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, str):
        return value if value else '""'
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return repr(value)
    raise SourceVerificationReportValueError(f"不支持序列化的结果值类型：{type(value).__name__}")


def _type_name(value: Any) -> str:
    return "NoneType" if value is None else type(value).__name__


def _header(sheet: Worksheet, values: tuple[str, ...]) -> None:
    _append(sheet, values)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"


def _append(sheet: Worksheet, values: Iterable[Any]) -> None:
    if sheet.max_row >= MAX_SHEET_ROWS and sheet.cell(sheet.max_row, 1).value is not None:
        raise SourceVerificationReportCapacityError(f"Sheet {sheet.title} 超过 {MAX_SHEET_ROWS} 行")
    row_number = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column_number, value in enumerate(values, start=1):
        cell = sheet.cell(row_number, column_number)
        if value is None:
            continue
        if isinstance(value, Path):
            value = str(value)
        if isinstance(value, str):
            _check_text(value)
            cell.value = value
            cell.data_type = "s"
        elif isinstance(value, (bool, int)):
            cell.value = value
        else:
            raise SourceVerificationReportValueError(f"不支持的报告单元格类型：{type(value).__name__}")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    if row_number > 1 and sheet.title == "差异明细":
        fill = _STATUS_FILLS.get(sheet.cell(row_number, 4).value)
        if fill:
            sheet.cell(row_number, 3).fill = PatternFill("solid", fgColor=fill)


def _check_text(value: str) -> None:
    if len(value) > MAX_CELL_TEXT:
        raise SourceVerificationReportValueError(f"文本长度 {len(value)} 超过 Excel 单元格上限 {MAX_CELL_TEXT}")
    if ILLEGAL_CHARACTERS_RE.search(value):
        raise SourceVerificationReportValueError("文本包含 Excel 不支持的控制字符")


def _widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _filter(sheet: Worksheet) -> None:
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions
