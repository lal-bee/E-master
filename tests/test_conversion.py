import io
import json
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, to_excel

from excel_qc import WorkbookError, convert_workbook
from excel_qc.converter import NS, _write_cells
from excel_qc.rules import (DATE_FIELDS, MappingError, YearOnlyError, load_mapping,
                            normalize_date, normalize_device_code)
from tests.fixtures import HEADERS, make_mapping, make_workbook, verify_download


@pytest.fixture
def mapping(tmp_path):
    return make_mapping(tmp_path / "映射.xlsx")


def single(tmp_path, fields, row):
    source = tmp_path / "原始台账.xlsx"
    book = Workbook()
    book.active.append(fields)
    book.active.append(row)
    book.save(source)
    book.close()
    return source


def output_book(result):
    assert result.success
    return load_workbook(io.BytesIO(result.output))


@pytest.mark.parametrize("value,expected", [
    ("2026-10-14", "2026-10-14"), ("2026/10/14", "2026-10-14"),
    ("2026.10.14", "2026-10-14"), ("2026-08", "2026-08-01"),
    ("2026/08", "2026-08-01"), ("2026.08", "2026-08-01"),
    (2026.08, "2026-08-01"), (date(2026, 10, 14), "2026-10-14"),
    (datetime(2026, 10, 14, 8, 30), "2026-10-14"),
    ("2026-10-14T09:30:00", "2026-10-14"), (None, None), ("", None),
    ("  ", None), ("非法日期", None), ("2026-13-01", None),
    ("2026.13", None), ("2026-02-30", None), ("2026-02/01", None),
])
def test_date_rules(value, expected):
    assert normalize_date(value, numeric=isinstance(value, (int, float))) == expected


@pytest.mark.parametrize("value", ["2026", " 2026 ", 2026, 2026.0, "2026.00"])
def test_year_only_is_blocking(value):
    with pytest.raises(YearOnlyError):
        normalize_date(value, numeric=isinstance(value, (int, float)))


@pytest.mark.parametrize("field", DATE_FIELDS)
@pytest.mark.parametrize("value,expected", [
    ("2026/08", "2026-08-01"), (2026.08, "2026-08-01"),
    (datetime(2026, 10, 14), "2026-10-14"), ("非法日期", None),
])
def test_all_three_date_columns(tmp_path, mapping, field, value, expected):
    result = convert_workbook(single(tmp_path, [field], [value]), mapping)
    book = output_book(result)
    assert book.active["A2"].value == expected
    assert result.dates_cleared == (expected is None)
    assert result.dates_converted == (expected is not None)
    book.close()


@pytest.mark.parametrize("epoch", [None, MAC_EPOCH])
def test_excel_serial_and_date_system(tmp_path, mapping, epoch):
    source = single(tmp_path, [DATE_FIELDS[0]], [0])
    book = load_workbook(source)
    if epoch:
        book.epoch = epoch
    book.active["A2"] = to_excel(datetime(2026, 10, 14), book.epoch)
    book.save(source)
    book.close()
    result = convert_workbook(source, mapping)
    out = output_book(result)
    assert out.active["A2"].value == "2026-10-14"
    out.close()


def test_numeric_october_with_two_decimal_places(tmp_path, mapping):
    source = single(tmp_path, [DATE_FIELDS[0]], [2026.1])
    book = load_workbook(source)
    book.active["A2"].number_format = "0.00"
    book.save(source)
    book.close()
    out = output_book(convert_workbook(source, mapping))
    assert out.active["A2"].value == "2026-10-01"
    out.close()


@pytest.mark.parametrize("value,expected", [
    ("ABC001", "ABC001"), ("000123", "000123"),
    ("000123456789012345678901234567890", "000123456789012345678901234567890"),
    ('A"B\\C', 'A"B\\C'), (None, None), ("", None), (" ", None),
    ("/", None), (" / ", None), ("null", None), ("NULL", None),
    ('{"注册代码":"ABC001"}', "ABC001"), ('{"注册代码":"/"}', None),
])
def test_device_code_rules(value, expected):
    actual = normalize_device_code(value)
    assert actual is None if expected is None else json.loads(actual) == {"注册代码": expected}


@pytest.mark.parametrize("value", [None, "", "   ", "/", "null"])
def test_empty_code_is_really_empty_cell(tmp_path, mapping, value):
    source = single(tmp_path, ["设备代码", "保留字段"], [value, "数据"])
    out = output_book(convert_workbook(source, mapping))
    assert out.active["A2"].value is None
    out.close()


def test_numeric_codes_never_float_or_scientific(tmp_path, mapping):
    source = single(tmp_path, ["设备代码"], [7])
    book = load_workbook(source)
    book.active["A2"].number_format = "000000"
    book.save(source)
    book.close()
    out = output_book(convert_workbook(source, mapping))
    assert json.loads(out.active["A2"].value)["注册代码"] == "000007"
    out.close()
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b"<v>7</v>", b"<v>1.234567890123456789E+20</v>")
    data = io.BytesIO()
    with ZipFile(data, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    out = output_book(convert_workbook(data.getvalue(), mapping))
    assert json.loads(out.active["A2"].value)["注册代码"] == "123456789012345678900"
    out.close()


def test_exact_mapping(tmp_path, mapping):
    source = single(tmp_path, ["设备小类名称", "设备小类编码"], ["离心泵", "旧编码"])
    result = convert_workbook(source, mapping)
    out = output_book(result)
    assert out.active["B2"].value == "D-LXB"
    assert out.active["A2"].value == "离心泵"
    assert result.subcategory_filled == 1
    out.close()


@pytest.mark.parametrize("value", ["离心", " 离心泵", "离心泵 ", "未知设备"])
def test_no_fuzzy_mapping(tmp_path, mapping, value):
    result = convert_workbook(single(tmp_path, ["设备小类名称", "设备小类编码"], [value, None]), mapping)
    assert not result.success and result.output is None
    assert result.issues[0].value == value and result.issues[0].row == 2
    assert result.issues[0].kind == "UNMATCHED"


def test_all_issues_collected_before_any_output(tmp_path, mapping, monkeypatch):
    source = make_workbook(tmp_path / "全部阻断.xlsx", bad=True)
    before = source.read_bytes()
    def forbidden(*args):
        pytest.fail("阻断时不能执行 XML 转换")
    monkeypatch.setattr("excel_qc.converter._write_cells", forbidden)
    result = convert_workbook(source, mapping)
    assert result.output is None
    assert len(result.issues) == 7
    assert len([i for i in result.issues if i.kind == "UNMATCHED"]) == 3
    assert len([i for i in result.issues if i.kind == "YEAR_ONLY"]) == 4
    assert {i.sheet for i in result.issues} == {"设备台账", "另一业务表"}
    assert all(i.row >= 2 and i.field and i.cell for i in result.issues)
    assert result.dates_converted == 0 and result.subcategory_filled == 0
    output = tmp_path / "不应生成.xlsx"
    with pytest.raises(WorkbookError):
        result.save(output)
    assert not output.exists()
    assert source.read_bytes() == before


def test_preserve_full_workbook_and_other_cells(tmp_path, mapping):
    source = make_workbook(tmp_path / "原件.xlsx")
    with ZipFile(source, "a") as archive:
        archive.writestr("customXml/preserved.xml", b"<unknown>keep exactly</unknown>")
        archive.writestr("xl/media/original.bin", b"unchanged image data")
        archive.comment = b"original zip comment"
    before = sha256(source.read_bytes()).hexdigest()
    result = convert_workbook(source, mapping)
    output = result.save(tmp_path / "结果.xlsx", source=source)
    assert sha256(source.read_bytes()).hexdigest() == before
    verify_download(output)
    assert result.sheets == ["设备台账"] and len(result.skipped_sheets) == 4
    assert (result.rows, result.subcategory_filled, result.dates_converted, result.dates_cleared, result.codes_formatted) == (3, 2, 5, 1, 1)
    with ZipFile(source) as a, ZipFile(output) as b:
        assert a.namelist() == b.namelist()
        assert a.comment == b.comment
        for name in a.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert a.read(name) == b.read(name), name
        original = ET.fromstring(a.read("xl/worksheets/sheet1.xml"))
        converted = ET.fromstring(b.read("xl/worksheets/sheet1.xml"))
        originals = {c.get("r"): c for c in original.findall(f".//{{{NS}}}c")}
        outputs = {c.get("r"): c for c in converted.findall(f".//{{{NS}}}c")}
        for coordinate, cell in originals.items():
            if coordinate.endswith("1") or coordinate.startswith(("A", "B", "H", "I")):
                assert ET.tostring(cell) == ET.tostring(outputs[coordinate]), coordinate
        for name in ["cols", "mergeCells", "dataValidations", "sheetViews", "tableParts", "hyperlinks"]:
            for a_node, b_node in zip(original.findall(f"{{{NS}}}{name}"), converted.findall(f"{{{NS}}}{name}")):
                assert ET.tostring(a_node) == ET.tostring(b_node)


def test_source_and_existing_output_cannot_be_overwritten(tmp_path, mapping):
    source = single(tmp_path, ["设备代码"], ["ABC"])
    result = convert_workbook(source, mapping)
    with pytest.raises(WorkbookError):
        result.save(source, source=source)
    with pytest.raises(FileExistsError):
        result.save(source)


def test_header_never_transformed_and_rows_not_reordered(tmp_path, mapping):
    source = single(tmp_path, ["设备代码", "序号"], ["A", 2])
    book = load_workbook(source)
    book.active.append(["B", 1])
    book.save(source)
    book.close()
    out = output_book(convert_workbook(source, mapping))
    assert list(out.active.values) == [("设备代码", "序号"), ('{"注册代码":"A"}', 2), ('{"注册代码":"B"}', 1)]
    out.close()


def test_multiple_business_sheets(tmp_path, mapping):
    source = single(tmp_path, ["设备代码"], ["ABC"])
    book = load_workbook(source)
    second = book.create_sheet("第二业务表")
    second.append(["设备制造日期"])
    second.append(["2026.08"])
    book.save(source)
    book.close()
    result = convert_workbook(source, mapping)
    assert len(result.sheets) == 2 and result.rows == 2
    out = output_book(result)
    assert out["第二业务表"]["A2"].value == "2026-08-01"
    out.close()


@pytest.mark.parametrize("headers,row", [
    (["设备小类名称"], ["离心泵"]),
    (["设备代码", "设备代码"], ["A", "B"]),
    (["设备代码"], ["=1+1"]),
])
def test_unsafe_structure_is_explained(tmp_path, mapping, headers, row):
    result = convert_workbook(single(tmp_path, headers, row), mapping)
    assert result.output is None and any(i.kind == "STRUCTURE" for i in result.issues)


def test_blank_categories_and_empty_dates(tmp_path, mapping):
    source = single(tmp_path, HEADERS, ["资产", None, None, None, "  ", None, None, "保留", None])
    result = convert_workbook(source, mapping)
    assert result.success and result.subcategory_filled == 0 and result.dates_cleared == 0


def test_no_business_sheet(tmp_path, mapping):
    source = single(tmp_path, ["辅助说明"], ["2026"])
    with pytest.raises(WorkbookError, match="没有找到"):
        convert_workbook(source, mapping)


def test_invalid_xlsx(mapping):
    with pytest.raises(WorkbookError, match="有效"):
        convert_workbook(b"not an excel file", mapping)


def test_mapping_file_and_conflicts(tmp_path):
    existing = load_mapping()
    assert existing["离心泵"] == "D-LXB" and len(existing) == 11
    conflict = make_mapping(tmp_path / "冲突.xlsx", [("离心泵", "A"), ("离心泵", "B"), ("静设备", "A")])
    with pytest.raises(MappingError) as error:
        load_mapping(conflict)
    assert "多个编码" in str(error.value) and "多个名称" in str(error.value)


def test_missing_mapping(tmp_path):
    with pytest.raises(MappingError):
        load_mapping(tmp_path / "不存在.xlsx")


def test_mapping_is_reloaded_each_execution(tmp_path):
    mapping = make_mapping(tmp_path / "映射.xlsx")
    source = single(tmp_path, ["设备小类名称", "设备小类编码"], ["新增小类", None])
    assert not convert_workbook(source, mapping).success
    make_mapping(mapping, [("新增小类", "NEW")])
    assert convert_workbook(source, mapping).success


@pytest.mark.parametrize("value", [True, False, "#VALUE!"])
def test_boolean_and_excel_error_dates_are_cleared(tmp_path, mapping, value):
    source = single(tmp_path, ["出厂日期"], [value])
    result = convert_workbook(source, mapping)
    out = output_book(result)
    assert out.active["A2"].value is None
    assert result.dates_cleared == 1 and result.dates_converted == 0
    out.close()


def test_mac_epoch_zero_is_a_valid_date(tmp_path, mapping):
    source = single(tmp_path, ["启用日期"], [0])
    book = load_workbook(source)
    book.epoch = MAC_EPOCH
    book.save(source)
    book.close()
    out = output_book(convert_workbook(source, mapping))
    assert out.active["A2"].value == "1904-01-01"
    out.close()


def test_code_without_name_has_source_coordinate(tmp_path, mapping):
    source = single(tmp_path, ["设备小类编码"], ["D-LXB"])
    result = convert_workbook(source, mapping)
    assert result.output is None
    assert [(i.kind, i.cell, i.value) for i in result.issues] == [("UNMATCHED", "A2", "D-LXB")]


def test_target_formula_reports_formula_text(tmp_path, mapping):
    result = convert_workbook(single(tmp_path, ["设备代码"], ["=1+1"]), mapping)
    assert result.issues[0].value == "=1+1"


def test_merged_target_blocks_without_losing_other_issues(tmp_path, mapping):
    source = single(tmp_path, ["设备代码", "出厂日期"], ["ABC", "2026"])
    book = load_workbook(source)
    book.active.merge_cells("A2:A3")
    book.save(source)
    book.close()
    result = convert_workbook(source, mapping)
    assert result.output is None
    assert {(i.kind, i.cell) for i in result.issues} == {("STRUCTURE", "A2"), ("YEAR_ONLY", "B2")}


def test_array_formula_from_other_column_cannot_overwrite_target(tmp_path, mapping):
    from openpyxl.worksheet.formula import ArrayFormula
    source = single(tmp_path, ["计算列", "设备代码"], [None, "ABC"])
    book = load_workbook(source)
    book.active["A2"] = ArrayFormula(ref="A2:B2", text="=A4:B4")
    book.save(source)
    book.close()
    result = convert_workbook(source, mapping)
    assert result.output is None
    assert result.issues[0].cell == "B2" and result.issues[0].value == "=A4:B4"


def test_value_precedes_cell_extensions_and_namespaces_survive():
    xml = f'''<s:worksheet xmlns:s="{NS}" xmlns:custom="urn:custom">
      <s:sheetData><s:row r="2"><s:c r="A2" s="7"><s:v>12</s:v>
        <s:extLst><s:ext uri="keep"><custom:data value="same"/></s:ext></s:extLst>
      </s:c></s:row></s:sheetData></s:worksheet>'''.encode()
    out = ET.fromstring(_write_cells(xml, {"A2": 'A"B<&'}))
    cell = out.find(f".//{{{NS}}}c")
    assert [n.tag for n in cell] == [f"{{{NS}}}is", f"{{{NS}}}extLst"]
    assert cell.get("s") == "7"
    assert cell.find(f"{{{NS}}}is/{{{NS}}}t").text == 'A"B<&'
    assert cell.find(".//{urn:custom}data").get("value") == "same"
