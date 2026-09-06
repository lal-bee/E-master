"""在测试中动态生成 Excel/CSV 样例，避免把二进制文件提交进仓库。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook


def build_workbook(path: Path, sheets: dict[str, list[list[Any]]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for sheet_name, rows in sheets.items():
        worksheet = book.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(list(row))
    book.save(path)
    return path


def build_trailing_note_workbook(path: Path) -> Path:
    """含标题行、表头、数据行（含重复行与稀疏数据行）及数据区末尾说明行的混乱样例。"""
    return build_workbook(
        path,
        {
            "设备台账": [
                ["2024年设备主数据台账", "", "", "", ""],
                ["设备编号", "设备名称", "规格型号", "启用日期", "备注"],
                ["EQ-001", "空压机", "GA-75", "2024-01-05", "进口"],
                ["EQ-002", "冷水机", "LS-100", "2024/2/10", "国产"],
                ["EQ-002", "冷水机", "LS-100", "2024/2/10", "国产"],
                ["", "备用机", "", "", ""],
                ["说明：本台账由设备管理部提供，字段口径以 2024 版台账为准", "", "", "", ""],
            ]
        },
    )


def build_merged_workbook(path: Path) -> Path:
    """含跨列合并标题单元格、表头与数据的样例。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    worksheet = book.active
    worksheet.title = "设备台账"
    worksheet.append(["2024年设备主数据台账"])
    worksheet.merge_cells("A1:E1")
    worksheet.append(["设备编号", "设备名称", "规格型号", "启用日期", "备注"])
    worksheet.append(["EQ-001", "空压机", "GA-75", "2024-01-05", "进口"])
    worksheet.append(["EQ-002", "冷水机", "LS-100", "2024/2/10", "国产"])
    book.save(path)
    return path


def build_xlsm_workbook(path: Path) -> Path:
    """生成最小 .xlsm 样例（openpyxl 可保存为可读取的 xlsm 包，不含真实 VBA 宏）。"""
    return build_workbook(
        path,
        {
            "设备台账": [
                ["设备编号", "设备名称"],
                ["EQ-001", "空压机"],
            ]
        },
    )


def build_mixed_type_workbook(path: Path) -> Path:
    """同一列混合数字、文本与空值的样例。"""
    return build_workbook(
        path,
        {
            "Sheet1": [
                ["设备编号", "设备名称", "数量"],
                [1, "空压机", 10],
                ["EQ-002", "", "3.5"],
                ["", "冷水机", ""],
                ["3.5", "备用机", 0],
            ]
        },
    )


def build_deep_header_workbook(path: Path) -> Path:
    """前 16 行是说明文字、真实表头位于第 17 行的样例（超出 15 行扫描窗口）。"""
    rows: list[list[Any]] = [[f"表前说明文字 {index}"] for index in range(1, 17)]
    rows.extend(
        [
            ["设备编号", "设备名称", "规格型号", "启用日期", "备注"],
            ["EQ-001", "空压机", "GA-75", "2024-01-05", "进口"],
            ["EQ-002", "冷水机", "LS-100", "2024/2/10", "国产"],
        ]
    )
    return build_workbook(path, {"设备台账": rows})


def write_csv(
    path: Path,
    rows: list[list[Any]],
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = "\r\n".join(delimiter.join(str(cell) for cell in row) for row in rows)
    path.write_text(payload, encoding=encoding, newline="")
    return path
