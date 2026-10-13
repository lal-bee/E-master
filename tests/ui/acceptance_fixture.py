"""P1-12 浏览器验收合成样本及实际下载回读；不进入业务运行链路。"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT_TEXT = "<img src=x onerror=alert(7)>"
SHEET = "<svg onload=alert(7)>"
LABEL = "label<em>名称</em>"
LONG_KEY = "000123456789012345678901234567890"


def create(root):
    root.mkdir(parents=True, exist_ok=True)

    def excel(name, headers, rows):
        path = root / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET
        for row in (headers, *rows):
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        workbook.save(path)
        workbook.close()
        return path

    source_header = ["源编号", "源名称" + SCRIPT_TEXT, "源计数", "源启用", "源长数值", "源数量"]
    system_header = ["系统编号", "系统名称", "系统计数", "系统启用", "系统长数值", "系统数量"]
    source = excel("原始源 & '演示'.xlsx", source_header, [
        ["0007", "=1+1", 0, False, 9876543210123456, "bad" + SCRIPT_TEXT],
        [LONG_KEY, SCRIPT_TEXT, 0, False, 17, 3],
        ["FORMULA_MINUS", "-SUM(A1)", 0, False, 17, 3],
    ])
    system = excel("系统导出 & '演示'.xlsx", system_header, [
        ["0007", "+SUM(A1)", 1, True, 9876543210123454, 2],
        [LONG_KEY, "@SUM(A1)", 1, True, 18, 3],
        ["FORMULA_MINUS", "=2+2", 1, True, 18, 3],
    ])
    names = ["record_id", LABEL, "count", "flag", "long_number", "quantity"]

    def config(role, header):
        return {
            "sheet_name": SHEET, "header_row_number": 1,
            "standard_fields": [{"name": name, "aliases": [alias]} for name, alias in zip(names, header)],
            "primary_key_fields": ["record_id"],
            "cleaning_rules": [{"standard_field":"quantity", "kind":"number", "rule_code":role + SCRIPT_TEXT}],
        }

    files = {"source":source, "system":system}
    for role, data in (
        ("original_config", config("SOURCE", source_header)),
        ("system_config", config("SYSTEM", system_header)),
        ("verification", {"key_fields":["record_id"], "comparison_fields":names[1:]}),
    ):
        path = root / (role + ".json")
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        files[role] = path
    invalid = root / "错误核验配置.json"
    invalid.write_text('{"key_fields":["record_id"],"comparison_fields":["not_in_schema"]}', encoding="utf-8")
    paths = {role:str(path) for role, path in files.items()}
    hashes = {str(path):hashlib.sha256(path.read_bytes()).hexdigest() for path in files.values()}
    return {"verify":paths, "invalid_config":str(invalid), "hashes":hashes}


def value_text(value):
    if value is None:
        return "None"
    if isinstance(value, dict):
        return value["value"]
    if isinstance(value, bool):
        return str(value)
    if value == "":
        return '""'
    return str(value)


def check_report(report, state, generic=False):
    """以实际 UI 结果为基准核对四 Sheet，不重算领域结论。"""
    result = json.loads(state.read_text(encoding="utf-8"))
    summary = result["summary"]
    book = load_workbook(report)
    try:
        assert book.sheetnames == ["核验概要", "差异统计", "差异明细", "异常记录"]
        overview = {row[0]:row[1] for row in book["核验概要"].iter_rows(min_row=2, values_only=True)}
        assert overview["原始源文件"] == result["original_source"]["source_file"]
        assert overview["系统导出文件"] == result["system_export"]["source_file"]
        assert overview["核验完成状态代码"] == str(summary["verification_complete"])
        assert overview["data_consistent"] == str(summary["data_consistent"])
        conclusion = "无法判定" if not summary["verification_complete"] else "一致" if summary["data_consistent"] else "不一致"
        assert overview["数据一致性结论"] == conclusion
        codes = {
            "ORIGINAL_SOURCE_INPUT":"original_source_input_row_count", "SYSTEM_EXPORT_INPUT":"system_export_input_row_count",
            "MATCHED_PAIRS":"matched_row_count", "MATCH":"unchanged_row_count", "FIELD_CHANGED":"changed_row_count",
            "MISSING_IN_SYSTEM":"missing_in_system_row_count", "EXTRA_IN_SYSTEM":"extra_in_system_row_count",
            "DUPLICATE_KEY":"duplicate_key_group_count", "UNVERIFIABLE":"unverifiable_row_count",
            "COMPARED_FIELDS":"compared_field_count", "EQUAL_FIELDS":"equal_field_count",
            "CHANGED_FIELDS":"changed_field_count", "UNVERIFIABLE_FIELDS":"unverifiable_field_count", "ISSUES":"issue_count",
        }
        statistics = {row[1]:row[2] for row in book["差异统计"].iter_rows(min_row=2, values_only=True)}
        assert all(statistics[code] == summary[field] for code, field in codes.items())
        assert all(statistics[code] == count for code, count in summary["issue_counts"])

        def table(name):
            rows = list(book[name].values)
            return [dict(zip(rows[0], row)) for row in rows[1:]]

        details = table("差异明细")
        expected = []
        for record in result["records"]:
            if record["status"] == "MATCH":
                continue
            visible = [diff for diff in record["field_differences"] if diff["status"] != "EQUAL"]
            expected.extend((record, diff) for diff in visible) if visible else expected.append((record, None))
        assert len(details) == len(expected)
        for actual, (record, difference) in zip(details, expected):
            assert actual["结果码"] == record["status"]
            assert [item["value"] for item in json.loads(actual["主键值（含类型）"])] == [value_text(value) for value in record["key_values"]]
            if difference is None:
                assert actual["明细层级"] == "行"
                continue
            assert actual["标准字段"] == difference["standard_field_name"]
            assert actual["字段结果码"] == difference["status"]
            assert actual["原因"] == difference["reason"]
            for prefix, side in (("源", "original_source"), ("系统", "system_export")):
                field = difference[side]
                assert actual[prefix+"原始值"] == value_text(field["original_value"])
                assert actual[prefix+"标准化值"] == value_text(field["normalized_value"])
                assert actual[prefix+"单元格"] == field["location"]["cell_address"]
                assert actual[prefix+"文件"] == field["location"]["source_file"]
                assert actual[prefix+"Sheet"] == field["location"]["sheet_name"]
                assert json.loads(actual[prefix+"应用规则"]) == field["applied_rules"]
        if not generic:
            assert {row["字段结果码"] for row in details} >= {"CHANGED", "UNVERIFIABLE"}
            assert any(json.loads(row["主键值（含类型）"])[0]["value"] == "0007" for row in details)
            assert any(json.loads(row["主键值（含类型）"])[0]["value"] == LONG_KEY for row in details)
            assert any(row["源原始值"] == "0" and row["源原始值类型"] == "int" for row in details)
            assert any(row["源原始值"] == "False" and row["源原始值类型"] == "bool" for row in details)
            texts = {row["源原始值"] for row in details} | {row["系统原始值"] for row in details}
            assert {"=1+1", "+SUM(A1)", "-SUM(A1)", "@SUM(A1)"} <= texts
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
        issues = table("异常记录")
        assert len(issues) == sum(max(1, len(issue["locations"])) for issue in result["issues"])
        assert {row["问题编码"] for row in issues} == {issue["code"] for issue in result["issues"]}
        expected_issues = [(issue, location, index if location else None) for issue in result["issues"]
                           for index, location in enumerate(issue["locations"] or [None], 1)]
        for actual, (issue, location, index) in zip(issues, expected_issues):
            assert actual["问题类别"] == issue["category"]
            assert actual["问题编码"] == issue["code"]
            assert actual["问题所属侧"] == issue["side"]
            assert actual["标准字段"] == issue["standard_field_name"]
            assert actual["原因"] == issue["message"]
            assert actual["阻断核验"] == ("是" if issue["blocks_verification"] else "否")
            assert actual["关联位置序号"] == index
            for column, key in (("文件","source_file"), ("Sheet","sheet_name"), ("物理行号","row_number"),
                                ("物理列号","column_number"), ("列字母","column_letter"),
                                ("单元格","cell_address"), ("原字段名","source_field_name")):
                assert actual[column] == (location[key] if location else None)
    finally:
        book.close()
    return {"report":"PASS", "statistics":"UI result matched", "details":len(details), "issues":len(issues), "formula_cells":0}


def check_template(path):
    book = load_workbook(path, read_only=True)
    try:
        assert book.sheetnames == ["演示数据"]
        assert list(book.active.values) == [
            ("编号", "名称", "类型编码"), ("0007", "水泵甲", "D-PUMP"), ("0008", "风机乙", "D-FAN"),
        ]
    finally:
        book.close()
    return {"template":"PASS", "rows":2}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--create", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--generic", action="store_true")
    parser.add_argument("--template", type=Path)
    args = parser.parse_args()
    result = create(args.create) if args.create else check_template(args.template) if args.template else check_report(args.report, args.result, args.generic)
    print(json.dumps(result, ensure_ascii=True))
