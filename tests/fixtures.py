"""只在指定临时目录生成合成测试 Excel，绝不修改业务原件。"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

HEADERS = ["资产编号", "设备小类名称", "设备小类编码", "出厂日期", "启用日期", "设备制造日期", "设备代码", "备注", "计算列"]


def make_mapping(path: Path, rows=None):
    book = Workbook()
    book.active.append(["设备小类名称", "设备小类编码"])
    for row in rows or [("离心泵", "D-LXB"), ("静设备", "JSB")]:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


def make_workbook(path: Path, *, bad=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "设备台账"
    sheet.append(HEADERS)
    sheet.append(["00001", "离心泵", "旧编码", "2026-10-14", "2026/08", 2026.08, "00012345678901234567890", "原样保留", "=1+1"])
    sheet.append(["00002", "静设备", None, datetime(2026, 10, 14, 9, 30), "无效日期", "2026.08", "/", "第二条", "=A3"])
    sheet.append(["00003", None, None, None, None, None, "   ", False, 100])
    if bad:
        sheet["B2"] = "不存在的小类一"
        sheet["B3"] = "不存在的小类二"
        sheet["D2"] = "2026"
        sheet["E3"] = 2025
        sheet["F4"] = "2024"
        second = book.create_sheet("另一业务表")
        second.append(["设备小类名称", "设备小类编码", "设备制造日期"])
        second.append(["第三个未知名称", None, "2023"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    sheet["C2"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet["D2"].number_format = "yyyy/mm/dd"
    sheet["A2"].comment = Comment("保留注释", "测试")
    sheet["H2"].hyperlink = "https://example.com/"
    sheet.freeze_panes = "C2"
    sheet.column_dimensions["A"].width = 22
    sheet.row_dimensions[2].height = 28
    sheet.row_dimensions[3].hidden = True
    validation = DataValidation(type="list", formula1='"离心泵,静设备"')
    sheet.add_data_validation(validation)
    validation.add("B2:B20")
    sheet.add_table(Table(displayName="EquipmentTable", ref="A1:I4"))
    auxiliary = book.create_sheet("辅助说明")
    auxiliary.append(["说明", "其他值", "其他"])
    auxiliary.append(["未知小类", "2026", "无效"])
    auxiliary.merge_cells("A4:C4")
    auxiliary["A4"] = "保留合并单元格"
    for title, state in [("隐藏辅助表", "hidden"), ("非常隐藏表", "veryHidden"), ("hidden@系统数据", "visible")]:
        helper = book.create_sheet(title)
        helper.append(["设备小类名称", "设备小类编码", "出厂日期"])
        helper.append(["不能匹配的辅助数据", None, "2026"])
        helper.sheet_state = state
    book.save(path)
    book.close()
    return path


def verify_download(path: Path):
    book = load_workbook(path, data_only=False)
    try:
        sheet = book["设备台账"]
        assert sheet["A2"].value == "00001"
        assert sheet["C2"].value == "D-LXB"
        assert sheet["C3"].value == "JSB"
        assert sheet["D2"].value == "2026-10-14"
        assert sheet["E2"].value == sheet["F2"].value == "2026-08-01"
        assert sheet["D3"].value == "2026-10-14"
        assert sheet["E3"].value is None
        assert sheet["F3"].value == "2026-08-01"
        assert sheet["G2"].value == '{"注册代码":"00012345678901234567890"}'
        assert sheet["G3"].value is None and sheet["G4"].value is None
        assert sheet["I2"].value == "=1+1" and sheet["I2"].data_type == "f"
        assert sheet["H4"].value is False
        assert sheet.freeze_panes == "C2"
        assert sheet["C2"].fill.fgColor.rgb == "00FFFF00"
        assert book["隐藏辅助表"].sheet_state == "hidden"
        assert book["非常隐藏表"].sheet_state == "veryHidden"
        return {"download_readback": "PASS", "dates": "PASS", "codes": "PASS", "preservation": "PASS"}
    finally:
        book.close()
