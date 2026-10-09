"""P1-10 报告专项：使用临时源/导出 Excel 与 openpyxl 回读报告。"""

from __future__ import annotations

from copy import deepcopy
from contextlib import closing
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import Workbook, load_workbook

from excel_qc import (
    CleaningKind,
    CleaningWorkbookConfig,
    FieldCleaningRule,
    FieldMappingConfig,
    OriginalSourceConfig,
    SourceVerificationConfig,
    StandardFieldDefinition,
    SystemExportConfig,
    verify_source_against_system_export,
)
from ui import export_source_verification_report
from ui.source_verification_report import (
    SourceVerificationReportCapacityError,
    SourceVerificationReportError,
    SourceVerificationReportValueError,
)
import ui.source_verification_report as report_module


def _excel(path: Path, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    for row in rows:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # 演示输入中的等号文本不作为 Excel 公式。
    workbook.save(path)
    return path


def _config(*, extra: bool = False, clean_day: bool = False) -> SourceVerificationConfig:
    source_fields = [StandardFieldDefinition("id", ("源编号",)), StandardFieldDefinition("name", ("源名称",))]
    system_fields = [StandardFieldDefinition("id", ("系统编号",)), StandardFieldDefinition("name", ("系统名称",))]
    comparison = ["name"]
    source_cleaning = CleaningWorkbookConfig()
    system_cleaning = CleaningWorkbookConfig()
    if extra:
        source_fields.append(StandardFieldDefinition("flag", ("源启用",)))
        system_fields.append(StandardFieldDefinition("flag", ("系统启用",)))
        comparison.append("flag")
    if clean_day:
        source_fields.append(StandardFieldDefinition("day", ("源日期",)))
        system_fields.append(StandardFieldDefinition("day", ("系统日期",)))
        comparison.append("day")
        source_cleaning = CleaningWorkbookConfig((FieldCleaningRule(
            "day", kind=CleaningKind.DATE, date_input_formats=("%Y-%m-%d",),
            date_output_format="%Y-%m-%d", rule_code="SOURCE_DATE",
        ),))
        system_cleaning = CleaningWorkbookConfig((FieldCleaningRule(
            "day", kind=CleaningKind.DATE, date_input_formats=("%Y-%m-%d",),
            date_output_format="%Y-%m-%d", rule_code="SYSTEM_DATE",
        ),))
    return SourceVerificationConfig(
        OriginalSourceConfig("数据", 1, FieldMappingConfig(tuple(source_fields)), source_cleaning),
        SystemExportConfig("数据", 1, FieldMappingConfig(tuple(system_fields)), system_cleaning, ("id",)),
        ("id",), tuple(comparison),
    )


def _result(tmp_path: Path, source_rows: list[list[object]], system_rows: list[list[object]], **config_options):
    source = _excel(tmp_path / "source.xlsx", source_rows)
    system = _excel(tmp_path / "system.xlsx", system_rows)
    before = (source.read_bytes(), system.read_bytes())
    result = verify_source_against_system_export(source, system, _config(**config_options))
    assert (source.read_bytes(), system.read_bytes()) == before
    return result


def _table(sheet) -> dict[str, object]:
    return {sheet.cell(row, 1).value: sheet.cell(row, 2).value for row in range(2, sheet.max_row + 1)}


def _rows(sheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]


def test_four_sheets_consistent_summary_style_and_empty_details(tmp_path: Path) -> None:
    result = _result(tmp_path, [["源编号", "源名称"], ["0007", "甲"]],
                     [["系统编号", "系统名称"], ["0007", "甲"]])
    before = deepcopy(result)
    report = export_source_verification_report(result, tmp_path / "report.xlsx")
    assert result == before
    with closing(load_workbook(report)) as workbook:
        assert workbook.sheetnames == ["核验概要", "差异统计", "差异明细", "异常记录"]
        overview = _table(workbook["核验概要"])
        assert overview["原始源文件"] == str(result.original_source.source_file)
        assert overview["系统导出文件"] == str(result.system_export.source_file)
        assert overview["核验完成状态"] == "已完成"
        assert overview["数据一致性结论"] == "一致"
        assert overview["核验时间"] == "P1-09 结果未提供"
        assert "T" in overview["报告生成时间"]
        assert workbook["差异明细"].max_row == workbook["异常记录"].max_row == 1
        for name in workbook.sheetnames:
            assert workbook[name]["A1"].font.bold
            assert workbook[name].freeze_panes == "A2"
        assert workbook["差异明细"]["A1"].value == "主键字段"
        assert workbook["异常记录"]["A1"].value == "问题类别"


def test_multiple_changed_fields_missing_extra_and_statistics(tmp_path: Path) -> None:
    result = _result(tmp_path,
                     [["源编号", "源名称", "源启用"], ["A", "旧", True], ["B", "仅源", False]],
                     [["系统编号", "系统名称", "系统启用"], ["A", "新", False], ["C", "仅系统", True]], extra=True)
    with closing(load_workbook(export_source_verification_report(result, tmp_path / "out.xlsx"))) as book:
        details = _rows(book["差异明细"])
        assert len(details) == 4
        changed = [row for row in details if row["结果码"] == "FIELD_CHANGED"]
        assert {row["标准字段"] for row in changed} == {"name", "flag"}
        assert all(row["明细层级"] == "字段" for row in changed)
        assert {row["结果码"] for row in details} == {"FIELD_CHANGED", "MISSING_IN_SYSTEM", "EXTRA_IN_SYSTEM"}
        assert all(row["标准字段"] is None and row["明细层级"] == "行" for row in details if row["结果码"] != "FIELD_CHANGED")
        assert book["差异明细"].auto_filter.ref
        stats = {row["状态码"]: (row["数量"], row["单位"]) for row in _rows(book["差异统计"])}
        assert stats["MATCHED_PAIRS"] == (1, "匹配对")
        assert stats["FIELD_CHANGED"] == (1, "结果记录")
        assert stats["CHANGED_FIELDS"] == (2, "字段项")
        assert stats["MISSING_IN_SYSTEM"][0] == stats["EXTRA_IN_SYSTEM"][0] == 1
        assert _table(book["核验概要"])["数据一致性结论"] == "不一致"


def test_unverifiable_duplicate_locations_and_confirmed_difference_coexist(tmp_path: Path) -> None:
    result = _result(tmp_path,
                     [["源编号", "源名称", "源日期"], ["D", "甲", "坏日期"],
                      ["D", "乙", "坏日期"], ["U", "旧", "坏日期"]],
                     [["系统编号", "系统名称", "系统日期"], ["D", "甲", "坏日期"],
                      ["U", "新", "坏日期"]], clean_day=True)
    with closing(load_workbook(export_source_verification_report(result, tmp_path / "report.xlsx"))) as book:
        overview = _table(book["核验概要"])
        assert overview["核验完成状态"] == "未完成"
        assert overview["数据一致性结论"] == "无法判定"
        details = _rows(book["差异明细"])
        assert any(row["结果码"] == "UNVERIFIABLE" and row["标准字段"] == "name" and row["字段结果码"] == "CHANGED" for row in details)
        assert any(row["结果码"] == "UNVERIFIABLE" and row["标准字段"] == "day" and row["字段结果码"] == "UNVERIFIABLE" for row in details)
        assert any(row["结果码"] == "DUPLICATE_KEY" and row["明细层级"] == "行" for row in details)
        issues = _rows(book["异常记录"])
        duplicate = [row for row in issues if row["问题编码"] == "DUPLICATE_KEY"]
        assert len(duplicate) == 3
        assert {(row["来源侧"], row["物理行号"]) for row in duplicate} == {
            ("ORIGINAL_SOURCE", 2), ("ORIGINAL_SOURCE", 3), ("SYSTEM_EXPORT", 2),
        }
        assert {row["问题所属侧"] for row in issues if row["问题编码"] == "CLEANING_FAILED"} == {
            "ORIGINAL_SOURCE", "SYSTEM_EXPORT",
        }
        assert _table(book["核验概要"])["原始源输入行数"] == 3
        assert _table(book["核验概要"])["系统导出输入行数"] == 2


def test_text_safety_codes_types_and_precise_coordinates(tmp_path: Path) -> None:
    long_code = "000123456789012345678901234567"
    result = _result(tmp_path,
                     [["源编号", "源名称", "源启用"], [long_code, "=HYPERLINK(\"x\")", False]],
                     [["系统编号", "系统名称", "系统启用"], [long_code, "@恶意", 0]], extra=True)
    with closing(load_workbook(export_source_verification_report(result, tmp_path / "report.xlsx"))) as book:
        details = _rows(book["差异明细"])
        assert len(details) == 2
        name_row = next(row for row in details if row["标准字段"] == "name")
        flag_row = next(row for row in details if row["标准字段"] == "flag")
        assert long_code in name_row["主键值（含类型）"]
        assert name_row["源原始值"] == '=HYPERLINK("x")'
        assert name_row["源文件"] == str(result.original_source.source_file)
        assert (name_row["源Sheet"], name_row["源行号"], name_row["源列号"], name_row["源单元格"]) == ("数据", 2, 2, "B2")
        assert (name_row["系统Sheet"], name_row["系统行号"], name_row["系统列号"], name_row["系统单元格"]) == ("数据", 2, 2, "B2")
        assert (flag_row["源原始值"], flag_row["源原始值类型"]) == ("False", "bool")
        assert (flag_row["系统原始值"], flag_row["系统原始值类型"]) == ("0", "int")
        header = [cell.value for cell in book["差异明细"][1]]
        column = header.index("源原始值") + 1
        cell = next(book["差异明细"].cell(row, column) for row in range(2, 4) if book["差异明细"].cell(row, column).value.startswith("="))
        assert cell.data_type == "s"
        assert book["差异明细"]["C2"].fill.patternType == "solid"


def test_none_empty_decimal_and_long_integer_serialization(tmp_path: Path) -> None:
    result = _result(tmp_path, [["源编号", "源名称"], ["A", "甲"]],
                     [["系统编号", "系统名称"], ["A", "乙"]])
    record = result.records[0]
    difference = record.field_differences[0]
    source_value = replace(difference.original_source, original_value=None, normalized_value="")
    system_value = replace(difference.system_export, original_value=Decimal("12.50"), normalized_value=123456789012345678901234567890)
    changed = replace(record, field_differences=(replace(difference, original_source=source_value, system_export=system_value),))
    synthetic = replace(result, records=(changed,))
    with closing(load_workbook(export_source_verification_report(synthetic, tmp_path / "values.xlsx"))) as book:
        row = _rows(book["差异明细"])[0]
        assert (row["源原始值"], row["源原始值类型"]) == ("None", "NoneType")
        assert (row["源标准化值"], row["源标准化值类型"]) == ('""', "str")
        assert (row["系统原始值"], row["系统原始值类型"]) == ("12.50", "Decimal")
        assert row["系统标准化值"] == "123456789012345678901234567890"
        assert row["系统标准化值类型"] == "int"


def test_file_level_issue_and_empty_input_do_not_become_pass(tmp_path: Path) -> None:
    result = _result(tmp_path, [["源编号", "源名称"]], [["系统编号", "系统名称"]])
    with closing(load_workbook(export_source_verification_report(result, tmp_path / "empty.xlsx"))) as book:
        assert _table(book["核验概要"])["数据一致性结论"] == "无法判定"
        assert book["差异明细"].max_row == 1
        issue = next(row for row in _rows(book["异常记录"]) if row["问题编码"] == "NO_DATA_ROWS")
        assert issue["物理行号"] is None and issue["文件"] is None


def test_missing_field_is_not_displayed_as_ordinary_empty_value(tmp_path: Path) -> None:
    result = _result(tmp_path, [["源编号"], ["A"]],
                     [["系统编号", "系统名称"], ["A", None]])
    with closing(load_workbook(export_source_verification_report(result, tmp_path / "missing.xlsx"))) as book:
        detail = _rows(book["差异明细"])[0]
        assert detail["结果码"] == "UNVERIFIABLE"
        assert detail["源字段状态"] == "FIELD_MISSING"
        assert detail["系统字段状态"] == "EMPTY"
        assert detail["源原始值"] == "None"
        assert any(row["问题编码"] == "FIELD_MISSING" for row in _rows(book["异常记录"]))


def test_overwrite_guards_input_aliases_and_existing_report(tmp_path: Path) -> None:
    result = _result(tmp_path, [["源编号", "源名称"], ["A", "甲"]],
                     [["系统编号", "系统名称"], ["A", "乙"]])
    source_before = result.original_source.source_file.read_bytes()
    system_before = result.system_export.source_file.read_bytes()
    with pytest.raises(SourceVerificationReportError, match="输入文件"):
        export_source_verification_report(result, result.original_source.source_file.parent / "." / "source.xlsx", overwrite=True)
    with pytest.raises(SourceVerificationReportError, match="输入文件"):
        export_source_verification_report(result, result.system_export.source_file, overwrite=True)
    assert (result.original_source.source_file.read_bytes(), result.system_export.source_file.read_bytes()) == (source_before, system_before)
    report = export_source_verification_report(result, tmp_path / "report.xlsx")
    before = report.read_bytes()
    with pytest.raises(FileExistsError):
        export_source_verification_report(result, report)
    assert report.read_bytes() == before
    assert export_source_verification_report(result, report, overwrite=True) == report
    with pytest.raises(SourceVerificationReportError, match=".xlsx"):
        export_source_verification_report(result, tmp_path / "wrong.csv")


def test_write_failure_keeps_existing_file_and_removes_temp(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    result = _result(tmp_path, [["源编号", "源名称"], ["A", "甲"]],
                     [["系统编号", "系统名称"], ["A", "乙"]])
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"original-report-content")
    existing = target.read_bytes()
    original_save = Workbook.save

    def failing_save(self, filename):
        Path(filename).write_bytes(b"partial")
        raise OSError("simulated write failure")

    monkeypatch.setattr(Workbook, "save", failing_save)
    with pytest.raises(OSError, match="simulated"):
        export_source_verification_report(result, target, overwrite=True)
    monkeypatch.setattr(Workbook, "save", original_save)
    assert target.read_bytes() == existing
    assert list(tmp_path.glob(".existing-*.xlsx")) == []


def test_invalid_character_and_capacity_fail_without_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    result = _result(tmp_path, [["源编号", "源名称"], ["A", "甲"]],
                     [["系统编号", "系统名称"], ["A", "乙"]])
    issue = replace(result.issues[0], message="非法\x01字符") if result.issues else None
    if issue is None:
        from excel_qc.source_verification import VerificationIssue, VerificationIssueCategory, VerificationSide
        issue = VerificationIssue(VerificationIssueCategory.INPUT, "X", "非法\x01字符", VerificationSide.BOTH)
    with pytest.raises(SourceVerificationReportValueError, match="控制字符"):
        export_source_verification_report(replace(result, issues=(issue,)), tmp_path / "invalid.xlsx")
    assert not (tmp_path / "invalid.xlsx").exists()
    assert list(tmp_path.glob(".invalid-*.xlsx")) == []
    monkeypatch.setattr(report_module, "MAX_SHEET_ROWS", 2)
    with pytest.raises(SourceVerificationReportCapacityError):
        export_source_verification_report(result, tmp_path / "capacity.xlsx")
    assert not (tmp_path / "capacity.xlsx").exists()


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_problem_text_formula_prefix_remains_text(tmp_path: Path, prefix: str) -> None:
    result = _result(tmp_path, [["源编号", "源名称"]], [["系统编号", "系统名称"]])
    issue = replace(result.issues[0], message=prefix + "危险内容")
    report = export_source_verification_report(replace(result, issues=(issue,)), tmp_path / "formula.xlsx")
    with closing(load_workbook(report)) as workbook:
        sheet = workbook["异常记录"]
        column = [cell.value for cell in sheet[1]].index("原因") + 1
        assert sheet.cell(2, column).value == prefix + "危险内容"
        assert sheet.cell(2, column).data_type == "s"


def test_overlong_text_fails_without_truncation(tmp_path: Path) -> None:
    result = _result(tmp_path, [["源编号", "源名称"]], [["系统编号", "系统名称"]])
    issue = replace(result.issues[0], message="甲" * 32768)
    with pytest.raises(SourceVerificationReportValueError, match="上限"):
        export_source_verification_report(replace(result, issues=(issue,)), tmp_path / "long.xlsx")
    assert not (tmp_path / "long.xlsx").exists()


def test_end_to_end_example_runs(tmp_path: Path) -> None:
    example = Path(__file__).resolve().parents[2] / "examples" / "source_verification_report_demo.py"
    completed = subprocess.run([sys.executable, str(example)], cwd=example.parents[1], text=True, capture_output=True, check=True)
    assert "核验概要" in completed.stdout
    assert "差异明细" in completed.stdout
    assert "异常记录" in completed.stdout
