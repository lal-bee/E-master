"""P1-08 系统导出 Excel 接入测试。"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import Workbook

from excel_qc import (
    CleaningKind,
    CleaningWorkbookConfig,
    FieldCleaningRule,
    FieldMappingConfig,
    FieldMappingStatus,
    StandardFieldDefinition,
    SystemExportConfig,
    SystemExportConfigError,
    SystemExportDataError,
    UnsupportedFileTypeError,
    ingest_system_export,
    load_system_export_config,
)
from excel_qc.loader import LoadedSheet, LoadedWorkbook
from excel_qc.profiler import WorkbookProfile
import excel_qc.system_export as system_export


def _workbook(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


def _config(
    *,
    sheet: str = "系统导出",
    header: int = 1,
    definitions: tuple[StandardFieldDefinition, ...] = (
        StandardFieldDefinition("record_id", ("编号", "ID")),
        StandardFieldDefinition("title", ("名称",)),
        StandardFieldDefinition("enabled", ("启用",)),
    ),
    keys: tuple[str, ...] = ("record_id",),
    cleaning: CleaningWorkbookConfig = CleaningWorkbookConfig(),
) -> SystemExportConfig:
    return SystemExportConfig(
        sheet_name=sheet,
        header_row_number=header,
        field_mapping=FieldMappingConfig(definitions),
        cleaning=cleaning,
        primary_key_fields=keys,
    )


def test_ingests_configured_sheet_header_mapping_and_cleaning(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "export.xlsx",
        {
            "说明": [["导出说明"]],
            "系统导出": [
                ["设备列表"],
                ["编号", "名称", "启用"],
                ["0012", "  泵　房 ", True],
            ],
        },
    )
    config = _config(
        header=2,
        cleaning=CleaningWorkbookConfig(
            (FieldCleaningRule("title", rule_code="TRIM_TITLE"),)
        ),
    )

    result = ingest_system_export(path, config)

    assert result.source_file == path
    assert result.sheet_name == "系统导出"
    assert result.header_row_number == 2
    assert result.rows[0].source_row_number == 3
    assert result.rows[0].values == {"record_id": "0012", "title": "泵 房", "enabled": True}
    assert result.rows[0].primary_key_value == "0012"
    title = next(field for field in result.rows[0].fields if field.standard_field_name == "title")
    assert (title.source_field_name, title.source_column_number, title.source_column_letter) == ("名称", 2, "B")
    assert title.original_value == "  泵　房 "
    assert title.normalized_value == "泵 房"
    assert title.applied_rules == ("TRIM_TITLE",)


def test_missing_fields_and_ambiguous_mapping_are_reported_without_guessing(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "ambiguous.xlsx",
        {"系统导出": [["编号", "别名", "别名"], ["A", "one", "two"]]},
    )
    config = _config(
        definitions=(
            StandardFieldDefinition("record_id", ("编号",)),
            StandardFieldDefinition("title", ("别名",)),
            StandardFieldDefinition("missing"),
        ),
        keys=("record_id",),
    )

    result = ingest_system_export(path, config)

    assert any(issue.code == "FIELD_MAPPING_CONFLICT" for issue in result.issues)
    assert any(issue.code == "MISSING_FIELD" and issue.standard_field_name == "title" for issue in result.issues)
    assert any(issue.code == "MISSING_FIELD" and issue.standard_field_name == "missing" for issue in result.issues)
    assert "title" not in result.rows[0].values
    assert len(result.rows) == 1


def test_same_source_alias_for_two_standard_fields_is_ambiguous(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "ambiguous-alias.xlsx",
        {"系统导出": [["编号", "名称"], ["R1", "one"]]},
    )
    config = _config(
        definitions=(
            StandardFieldDefinition("record_id", ("编号",)),
            StandardFieldDefinition("title", ("名称",)),
            StandardFieldDefinition("label", ("名称",)),
        ),
    )

    result = ingest_system_export(path, config)

    assert any(issue.code == "FIELD_MAPPING_AMBIGUOUS" for issue in result.issues)
    assert "title" not in result.rows[0].values
    assert "label" not in result.rows[0].values


def test_missing_sheet_is_reported_and_returns_consumable_empty_result(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "wrong-sheet.xlsx", {"Other": [["编号"], ["1"]]})

    result = ingest_system_export(path, _config())

    assert result.rows == ()
    assert result.issues[0].code == "SHEET_NOT_FOUND"
    assert result.summary.row_count == 0


def test_preserves_zero_false_blank_and_text_leading_zero(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "values.xlsx",
        {"系统导出": [["编号", "名称", "启用"], [0, None, False], ["0007", "x", 0]]},
    )

    result = ingest_system_export(path, _config())

    assert result.rows[0].values == {"record_id": 0, "title": None, "enabled": False}
    assert result.rows[0].primary_key_value == 0
    assert result.rows[1].values["record_id"] == "0007"
    assert result.rows[1].values["enabled"] == 0
    assert any(issue.code == "EMPTY_PRIMARY_KEY" for issue in result.issues) is False


def test_none_and_empty_string_remain_distinguishable_in_result(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "synthetic.xlsx"
    raw = LoadedWorkbook(
        path=path,
        file_type="xlsx",
        sheets=(
            LoadedSheet(
                "系统导出",
                (
                    ("编号", "名称", "启用"),
                    (None, "none row", False),
                    ("", "empty string row", False),
                ),
            ),
        ),
    )
    monkeypatch.setattr(system_export, "load_workbook", lambda _: raw)
    monkeypatch.setattr(
        system_export,
        "profile_workbook",
        lambda _: WorkbookProfile(path=path, sheets=()),
    )

    result = system_export.ingest_system_export(path, _config())

    assert result.rows[0].values["record_id"] is None
    assert result.rows[1].values["record_id"] == ""


def test_empty_duplicate_and_composite_keys_are_retained_and_reported(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "keys.xlsx",
        {
            "系统导出": [
                ["编号", "名称", "启用", "分区"],
                ["", "empty", True, "N"],
                ["A", "first", True, "N"],
                ["A", "duplicate", True, "N"],
                ["A", "different partition", True, "S"],
            ]
        },
    )
    config = _config(
        definitions=(
            StandardFieldDefinition("record_id", ("编号",)),
            StandardFieldDefinition("title", ("名称",)),
            StandardFieldDefinition("enabled", ("启用",)),
            StandardFieldDefinition("zone", ("分区",)),
        ),
        keys=("record_id", "zone"),
    )

    result = ingest_system_export(path, config)

    assert len(result.rows) == 4
    assert [row.primary_key_values for row in result.rows] == [
        (None, "N"), ("A", "N"), ("A", "N"), ("A", "S")
    ]
    assert [issue.code for issue in result.issues].count("EMPTY_PRIMARY_KEY") == 1
    duplicate = next(issue for issue in result.issues if issue.code == "DUPLICATE_PRIMARY_KEY")
    assert duplicate.row_number == 4
    assert "第 3 行" in duplicate.message


def test_composite_keys_do_not_collide_by_string_concatenation(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "non-colliding-keys.xlsx",
        {
            "系统导出": [
                ["编号", "名称", "启用", "分区"],
                ["ab", "first", True, "c"],
                ["a", "second", False, "bc"],
            ]
        },
    )
    config = _config(
        definitions=(
            StandardFieldDefinition("record_id", ("编号",)),
            StandardFieldDefinition("title", ("名称",)),
            StandardFieldDefinition("enabled", ("启用",)),
            StandardFieldDefinition("zone", ("分区",)),
        ),
        keys=("record_id", "zone"),
    )

    result = ingest_system_export(path, config)

    assert [row.primary_key_values for row in result.rows] == [("ab", "c"), ("a", "bc")]
    assert not any(issue.code == "DUPLICATE_PRIMARY_KEY" for issue in result.issues)


def test_cleaning_failure_keeps_raw_value_and_source_coordinates(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "bad-date.xlsx",
        {"系统导出": [["编号", "日期", "启用"], ["D1", "not-a-date", True]]},
    )
    config = _config(
        definitions=(
            StandardFieldDefinition("record_id", ("编号",)),
            StandardFieldDefinition("date", ("日期",)),
            StandardFieldDefinition("enabled", ("启用",)),
        ),
        cleaning=CleaningWorkbookConfig(
            (FieldCleaningRule("date", kind=CleaningKind.DATE, date_output_format="%Y-%m-%d", rule_code="DATE_ISO"),)
        ),
    )

    result = ingest_system_export(path, config)

    date_value = next(field for field in result.rows[0].fields if field.standard_field_name == "date")
    assert date_value.original_value == "not-a-date"
    assert date_value.normalized_value == "not-a-date"
    issue = next(issue for issue in result.issues if issue.code == "CLEANING_FAILED")
    assert (issue.row_number, issue.column_number, issue.column_letter) == (2, 2, "B")


def test_input_workbook_is_unchanged_and_empty_intermediate_rows_are_preserved(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "readonly.xlsx",
        {"系统导出": [["编号", "名称", "启用"], ["A", "first", True], [None, None, None], ["B", "last", False]]},
    )
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    result = ingest_system_export(path, _config())

    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    assert [row.source_row_number for row in result.rows] == [2, 3, 4]


def test_json_config_loader_and_downstream_result_shape(tmp_path: Path) -> None:
    config_path = tmp_path / "config.json"
    config_path.write_text(
        json.dumps(
            {
                "sheet_name": "系统导出",
                "header_row_number": 1,
                "standard_fields": [
                    {"name": "record_id", "aliases": ["编号"]},
                    {"name": "title", "aliases": ["名称"]},
                ],
                "cleaning_rules": [{"standard_field": "title", "kind": "text", "rule_code": "TITLE"}],
                "primary_key_fields": ["record_id"],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    config = load_system_export_config(config_path)
    path = _workbook(tmp_path / "consumer.xlsx", {"系统导出": [["编号", "名称"], ["R1", "Unit"]]})

    result = ingest_system_export(path, config)

    assert result.rows[0].values["record_id"] == "R1"
    assert result.rows[0].primary_key_fields == ("record_id",)
    assert result.mapping_result.summary.match_count == 2
    assert result.cleaning_result.summary.row_count == 1
    assert result.summary.row_count == len(result.rows)


def test_empty_sheet_header_only_and_no_effective_data_are_explicit(tmp_path: Path) -> None:
    empty_sheet_path = _workbook(tmp_path / "empty-sheet.xlsx", {"系统导出": []})
    with pytest.raises(SystemExportDataError, match="表头行 1 超出"):
        ingest_system_export(empty_sheet_path, _config())

    header_only = _workbook(
        tmp_path / "header-only.xlsx",
        {"系统导出": [["编号", "名称", "启用"]]},
    )
    result = ingest_system_export(header_only, _config())
    assert result.rows == ()
    assert result.summary.row_count == 0
    assert not any(issue.code in {"EMPTY_PRIMARY_KEY", "DUPLICATE_PRIMARY_KEY"} for issue in result.issues)

    no_effective_data = _workbook(
        tmp_path / "no-effective-data.xlsx",
        {"系统导出": [["编号", "名称", "启用"], [None, None, None]]},
    )
    result = ingest_system_export(no_effective_data, _config())
    assert result.rows == ()
    assert result.summary.row_count == 0


def test_non_xlsx_input_is_rejected_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "export.xls"
    path.write_bytes(b"unsupported input")

    with pytest.raises(UnsupportedFileTypeError, match="仅支持 .xlsx"):
        ingest_system_export(path, _config())


def test_temporary_excel_end_to_end_ingestion_integration(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "system-export-integration.xlsx",
        {
            "封面": [["某系统导出示例"]],
            "设备清单": [
                ["系统设备台账"],
                ["记录号", "设备描述", "状态"],
                ["00031", "  风机 A ", False],
                ["00032", "泵 B", True],
            ],
        },
    )
    config = SystemExportConfig(
        sheet_name="设备清单",
        header_row_number=2,
        field_mapping=FieldMappingConfig(
            (
                StandardFieldDefinition("record_id", ("记录号",)),
                StandardFieldDefinition("description", ("设备描述",)),
                StandardFieldDefinition("status", ("状态",)),
            )
        ),
        cleaning=CleaningWorkbookConfig(
            (FieldCleaningRule("description", rule_code="TRIM_DESCRIPTION"),)
        ),
        primary_key_fields=("record_id",),
    )

    result = ingest_system_export(path, config)

    assert result.summary.row_count == 2
    assert result.rows[0].source_row_number == 3
    assert result.rows[0].values == {
        "record_id": "00031",
        "description": "风机 A",
        "status": False,
    }
    assert result.rows[1].primary_key_value == "00032"
    assert not result.issues


def test_minimal_example_runs_with_temporary_demo_workbook(tmp_path: Path) -> None:
    repo_root = Path(__file__).resolve().parents[2]
    path = _workbook(
        tmp_path / "demo-export.xlsx",
        {
            "封面": [["演示数据"]],
            "设备清单": [
                ["系统设备台账"],
                ["记录号", "设备描述", "状态"],
                ["00042", "  example item ", "enabled"],
            ],
        },
    )

    completed = subprocess.run(
        [
            sys.executable,
            str(repo_root / "examples" / "system_export_demo.py"),
            str(path),
        ],
        cwd=repo_root,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "数据行: 1" in completed.stdout
    assert "主键='00042'" in completed.stdout
    assert "description" in completed.stdout


def test_invalid_config_and_unmatched_primary_key_are_rejected(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "bad-key.xlsx", {"系统导出": [["编号"], ["1"]]})
    with pytest.raises(SystemExportConfigError, match="主键字段未定义"):
        ingest_system_export(path, _config(keys=("unknown",)))

    config_path = tmp_path / "bad.json"
    config_path.write_text('{"sheet_name":"S","header_row_number":0,"standard_fields":[]}', encoding="utf-8")
    with pytest.raises(SystemExportConfigError, match="header_row_number"):
        load_system_export_config(config_path)
