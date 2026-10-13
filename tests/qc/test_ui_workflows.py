"""P1-12 编排、版本失效、门禁与输入只读的集成断言。"""
from __future__ import annotations

import json
from pathlib import Path
from threading import Event, Thread
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest
from openpyxl import load_workbook

from examples.ui_demo import create_demo
from ui import workflows
from ui.local_app import LocalServer
from ui.workflows import WorkflowError, WorkflowSession, display_value, user_message


@pytest.fixture
def demo(tmp_path):
    return create_demo(tmp_path / "demo")


def session(tmp_path, demo, flow="prepare", fail=False):
    result = WorkflowSession(tmp_path / "session" / flow)
    result.files = dict(demo[flow])
    if fail:
        result.files["source"] = demo["fail_source"]
    result.configure(demo["settings"] if flow == "prepare" else {})
    return result


def test_prepare_calls_each_public_stage_and_exports_only_consumed_result(tmp_path, demo, monkeypatch):
    names = ["profile_configured_sheet", "map_workbook_fields", "clean_workbook",
             "map_equipment_type_codes", "build_standard_template", "run_quality_gate"]
    calls = []
    for name in names:
        original = getattr(workflows, name)

        def spy(*args, _name=name, _original=original, **kwargs):
            calls.append(_name)
            return _original(*args, **kwargs)

        monkeypatch.setattr(workflows, name, spy)
    s = session(tmp_path, demo)
    inputs = {p: p.read_bytes() for p in s.files.values()}
    response = s.run("prepare", s.revision)
    assert calls == names
    assert response["exportable"]
    result = s.result
    assert result.profile.sheets[0].header_row == 2
    assert result.cleaning.sheet_results[0].rows[0].source_row_number == 3
    assert result.template.sheet_results[0].rows[0].values == ("0007", "水泵甲", "D-PUMP")
    output = s.export("prepare", s.revision)
    assert calls == names  # 导出不重新运行领域步骤
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames == ["演示数据"]
        assert list(workbook.active.values)[1] == ("0007", "水泵甲", "D-PUMP")
    finally:
        workbook.close()
    assert all(path.read_bytes() == data for path, data in inputs.items())


def test_fail_gate_keeps_details_and_never_calls_export(tmp_path, demo, monkeypatch):
    s = session(tmp_path, demo, fail=True)
    response = s.run("prepare", s.revision)
    assert not response["exportable"]
    assert s.result.gate.status.value == "FAIL"
    assert s.result.gate.issues
    monkeypatch.setattr(workflows, "export_standard_template", lambda *a: pytest.fail("FAIL 不得调用模板导出"))
    with pytest.raises(WorkflowError, match="质量门禁 FAIL"):
        s.export("prepare", s.revision)
    assert not (s.root / "exports").exists()


def test_verify_calls_engine_then_report_with_differences_and_unverifiable(tmp_path, demo, monkeypatch):
    s = session(tmp_path, demo, "verify")
    before = {p: p.read_bytes() for p in s.files.values()}
    calls = []
    engine = workflows.verify_source_against_system_export
    report = workflows.export_source_verification_report

    def verify(*args):
        calls.append("engine")
        return engine(*args)

    def export(result, output):
        calls.append("report")
        assert result is s.result
        return report(result, output)

    monkeypatch.setattr(workflows, "verify_source_against_system_export", verify)
    monkeypatch.setattr(workflows, "export_source_verification_report", export)
    assert s.run("verify", s.revision)["exportable"]
    assert s.result.summary.changed_field_count == 1
    assert s.result.summary.unverifiable_field_count == 1
    assert not s.result.summary.verification_complete
    mixed = next(r for r in s.result.records if r.key_values == ("CHANGED",))
    assert mixed.status.value == "UNVERIFIABLE"
    assert {d.status.value for d in mixed.field_differences} == {"CHANGED", "UNVERIFIABLE"}
    assert mixed.field_differences[0].original_source.location.cell_address == "B4"
    output = s.export("verify", s.revision)
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames == ["核验概要", "差异统计", "差异明细", "异常记录"]
    finally:
        workbook.close()
    assert calls == ["engine", "report"]
    assert all(p.read_bytes() == data for p, data in before.items())


@pytest.mark.parametrize("flow,role", [(flow, role) for flow, roles in workflows.ROLES.items() for role in roles])
def test_every_file_role_change_invalidates_old_result(tmp_path, demo, flow, role):
    s = session(tmp_path, demo, flow)
    s.run(flow, s.revision)
    revision = s.revision
    previous = s.files[role]
    s.upload(flow, role, previous.name, previous.read_bytes())
    assert s.revision > revision and s.result is None
    assert s.files[role] != previous
    assert previous.exists()  # 来源副本不因换输入而丢失
    with pytest.raises(WorkflowError, match="有效结果"):
        s.export(flow, revision)


@pytest.mark.parametrize("settings", [{"sheet":"空白Sheet","header":2}, {"sheet":"演示数据","header":1}])
def test_sheet_or_header_change_invalidates_result(tmp_path, demo, settings):
    s = session(tmp_path, demo)
    s.run("prepare", s.revision)
    s.configure(settings)
    assert s.result is None
    with pytest.raises(WorkflowError, match="有效结果"):
        s.export("prepare", s.revision)


@pytest.mark.parametrize("failure", ["missing", "invalid_config", "runtime"])
def test_run_failure_removes_previous_result_and_shows_no_traceback(tmp_path, demo, monkeypatch, failure):
    s = session(tmp_path, demo)
    s.run("prepare", s.revision)
    if failure == "missing":
        s.files["source"] = tmp_path / "missing.xlsx"
    elif failure == "invalid_config":
        s.files["mapping"].write_text('{"standard_fields": "wrong-type"}', encoding="utf-8")
    else:
        monkeypatch.setattr(workflows, "prepare_source", lambda *a: (_ for _ in ()).throw(RuntimeError("private stack details")))
    with pytest.raises(Exception) as caught:
        s.run("prepare", s.revision)
    assert s.result is None and "失败" in s.status
    message = user_message(caught.value)
    assert "Traceback" not in message and "private stack details" not in message
    with pytest.raises(WorkflowError):
        s.export("prepare", s.revision)


def test_missing_roles_and_invalid_json_are_locatable(tmp_path, demo):
    s = session(tmp_path, demo)
    s.files.pop("gate")
    with pytest.raises(WorkflowError, match="文件未选择完整"):
        s.run("prepare", s.revision)
    with pytest.raises(WorkflowError, match="第 2 行"):
        s.upload("prepare", "gate", "gate.json", b'{\n bad}')


def test_reject_stale_run_and_never_export_without_run(tmp_path, demo):
    s = session(tmp_path, demo)
    with pytest.raises(WorkflowError):
        s.export("prepare", s.revision)
    old = s.revision
    s.invalidate()
    with pytest.raises(WorkflowError, match="版本已改变"):
        s.run("prepare", old)


def test_bad_upload_invalidates_and_cannot_escape_input_directory(tmp_path, demo):
    s = session(tmp_path, demo)
    s.run("prepare", s.revision)
    with pytest.raises(WorkflowError, match="只接受"):
        s.upload("prepare", "source", "wrong.xls", b"bad")
    assert s.result is None and "source" not in s.files
    path = demo["prepare"]["source"]
    s.upload("prepare", "source", "../../outside.xlsx", path.read_bytes())
    assert s.files["source"].is_relative_to(s.root)
    assert s.files["source"].name == "outside.xlsx"


def test_verification_config_rejects_key_disagreement_and_unknown_rule(tmp_path, demo):
    s = session(tmp_path, demo, "verify")
    path = s.files["verification"]
    path.write_text(json.dumps({"key_fields":["different"],"comparison_fields":["label"]}), encoding="utf-8")
    with pytest.raises(WorkflowError, match="必须与核验"):
        s.run("verify", s.revision)
    path.write_text(json.dumps({"key_fields":["record_id"],"comparison_fields":["label"],"comparison_rules":[{"standard_field":"label","kind":"AI"}]}), encoding="utf-8")
    with pytest.raises(WorkflowError, match="kind 无效"):
        s.run("verify", s.revision)


@pytest.mark.parametrize("flow", ["prepare", "verify"])
def test_repeated_exports_create_new_files_and_never_replace_input(tmp_path, demo, flow, monkeypatch):
    s = session(tmp_path, demo, flow)
    s.run(flow, s.revision)
    first = s.export(flow, s.revision)
    first_bytes = first.read_bytes()
    second = s.export(flow, s.revision)
    assert first != second and first.read_bytes() == first_bytes
    # 即使文件身份碰撞，也在调用底层导出器前阻断输入覆盖。
    class FixedId:
        hex = "collision"
    monkeypatch.setattr(workflows, "uuid4", FixedId)
    protected = s.root / "exports" / f"{flow}-collision.xlsx"
    protected.write_bytes(s.files["source"].read_bytes())
    s.files["source"] = protected
    before = protected.read_bytes()
    with pytest.raises(WorkflowError, match="不得覆盖输入"):
        s.export(flow, s.revision)
    assert protected.read_bytes() == before


def test_display_preserves_types_and_nonfinite_values():
    assert display_value((None, "", 0, False)) == [None, "", 0, False]
    assert display_value(float("nan")) == {"type":"float", "value":"nan"}
    assert display_value(9007199254740993) == {"type":"int", "value":"9007199254740993"}


@pytest.fixture
def server(tmp_path):
    app = LocalServer(0, tmp_path / "http")
    thread = Thread(target=app.serve_forever, daemon=True)
    thread.start()
    yield app
    app.shutdown()
    app.server_close()
    thread.join(timeout=3)


def post(server, token, action, body, **headers):
    request = Request(f"http://127.0.0.1:{server.server_port}/api/{action}",
                      data=json.dumps(body).encode(),
                      headers={"Content-Type":"application/json", "X-Session-Token":token, **headers})
    return urlopen(request, timeout=10)


def test_http_both_flows_export_gate_and_invalidation(server):
    token = server.new_session()
    for flow in ("prepare", "verify"):
        with post(server, token, "demo", {"flow":flow}) as response:
            revision = json.load(response)["revision"]
        with post(server, token, "run", {"flow":flow,"revision":revision,"original_confirmed":True}) as response:
            assert json.load(response)["exportable"]
        with post(server, token, "export", {"flow":flow,"revision":revision}) as response:
            assert response.read()[:2] == b"PK"
        post(server, token, "invalidate", {"flow":flow}).close()
        with pytest.raises(HTTPError) as error:
            post(server, token, "export", {"flow":flow,"revision":revision})
        assert "有效结果" in json.load(error.value)["error"]
    with post(server, token, "demo", {"flow":"prepare","fail":True}) as response:
        revision = json.load(response)["revision"]
    with post(server, token, "run", {"flow":"prepare","revision":revision}) as response:
        assert not json.load(response)["exportable"]
    with pytest.raises(HTTPError) as error:
        post(server, token, "export", {"flow":"prepare","revision":revision})
    assert "FAIL" in json.load(error.value)["error"]


def test_http_protects_local_access_and_duplicate_actions(server, monkeypatch):
    assert server.server_address[0] == "127.0.0.1"
    token = server.new_session()
    with pytest.raises(HTTPError) as error:
        post(server, token, "demo", {"flow":"prepare"}, Origin="https://unrelated.example")
    assert error.value.code == 400
    with pytest.raises(HTTPError) as error:
        post(server, "wrong", "demo", {"flow":"prepare"})
    assert error.value.code == 403
    with post(server, token, "demo", {"flow":"prepare"}) as response:
        revision = json.load(response)["revision"]
    current = server.sessions[token]["prepare"]
    inputs = dict(current.files)
    started, finish = Event(), Event()
    original = workflows.prepare_source

    def delayed(*args):
        started.set()
        assert finish.wait(5)
        return original(*args)

    monkeypatch.setattr(workflows, "prepare_source", delayed)
    replies = []

    def run():
        with post(server, token, "run", {"flow":"prepare", "revision":revision}) as response:
            replies.append(json.load(response))

    thread = Thread(target=run)
    thread.start()
    assert started.wait(3)
    try:
        for action, extra in (("run", {"revision":revision}), ("invalidate", {"role":"source"}),
                              ("settings", {"settings":{"sheet":"new","header":1}}),
                              ("export", {"revision":revision})):
            with pytest.raises(HTTPError) as error:
                post(server, token, action, {"flow":"prepare", **extra})
            assert error.value.code == 409
        assert current.files == inputs and current.revision == revision
    finally:
        finish.set()
        thread.join(timeout=5)
    assert not thread.is_alive() and len(replies) == 1
    assert current.result_revision == revision
    assert replies[0]["result"]["profile"]["path"] == str(inputs["source"])


def test_browser_sessions_keep_distinct_inputs_results_and_downloads(server):
    from io import BytesIO
    tokens = [server.new_session(), server.new_session()]
    revisions = []
    for token, fail in zip(tokens, (False, True)):
        with post(server, token, "demo", {"flow":"prepare", "fail":fail}) as response:
            revision = json.load(response)["revision"]
        with post(server, token, "run", {"flow":"prepare", "revision":revision}) as response:
            assert json.load(response)["exportable"] is not fail
        revisions.append(revision)
    first, second = [server.sessions[token]["prepare"] for token in tokens]
    assert first.files["source"] != second.files["source"]
    assert first.result.gate.can_proceed and not second.result.gate.can_proceed
    with post(server, tokens[0], "export", {"flow":"prepare", "revision":revisions[0]}) as response:
        content = response.read()
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        assert list(workbook.active.values)[1] == ("0007", "水泵甲", "D-PUMP")
    finally:
        workbook.close()
    with pytest.raises(HTTPError):
        post(server, tokens[1], "export", {"flow":"prepare", "revision":revisions[1]})
    post(server, tokens[1], "invalidate", {"flow":"prepare","role":"source"}).close()
    assert first.result is not None and second.result is None
    assert first.files["source"].exists()


def test_http_page_assets_and_error_are_user_readable(server):
    root = f"http://127.0.0.1:{server.server_port}"
    with urlopen(root) as response:
        html = response.read().decode()
        assert "原始资料整理" in html and "导入后回源核验" in html
        assert "__SESSION_VALUE__" not in html
        assert 'const SESSION_TOKEN = "' in html
    token = server.new_session()
    with pytest.raises(HTTPError) as error:
        post(server, token, "run", {"flow":"verify","revision":0})
    assert "最初收集" in json.load(error.value)["error"]
    with urlopen(root + "/app.js") as response:
        assert "旧结果失效" in response.read().decode()


@pytest.mark.parametrize("failure", ["unreadable_file", "invalid_base64"])
def test_failed_file_change_clears_server_input_and_cannot_rerun_old_file(server, failure):
    token = server.new_session()
    with post(server, token, "demo", {"flow":"prepare"}) as response:
        revision = json.load(response)["revision"]
    post(server, token, "run", {"flow":"prepare", "revision":revision}).close()
    session = server.sessions[token]["prepare"]
    previous_source = session.files["source"]
    if failure == "unreadable_file":
        # 文件选择后先撤销该角色，再读取 FileReader；读失败也不能沿用旧源。
        post(server, token, "invalidate", {"flow":"prepare", "role":"source"}).close()
    else:
        with pytest.raises(HTTPError):
            post(server, token, "upload", {"flow":"prepare", "role":"source", "name":"new.xlsx", "data":"!bad-base64!"})
    assert session.result is None
    assert "source" not in session.files
    assert previous_source.exists()
    with pytest.raises(HTTPError) as error:
        post(server, token, "run", {"flow":"prepare", "revision":session.revision})
    assert "文件未选择完整" in json.load(error.value)["error"]
