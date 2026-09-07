"""Excel 核验报告导出（P0-10）测试。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_qc.models import (
    CellLocation,
    ValidationErrorType,
    ValidationIssue,
    ValidationResult,
    ValidationSource,
    ValidationSummary,
)
from tests.qc.helpers import build_workbook
from ui.excel_report import export_validation_report


def _issue(
    *,
    source_file: Path,
    sheet_name: str,
    row_number: int,
    column_number: int,
    column_letter: str,
    field_name: str,
    error_type: ValidationErrorType,
    error_code: str,
    actual_value: str,
    expected_value: str = "",
) -> ValidationIssue:
    return ValidationIssue(
        location=CellLocation(
            source_file=source_file,
            sheet_name=sheet_name,
            row_number=row_number,
            column_number=column_number,
            column_letter=column_letter,
            column_name=field_name,
        ),
        error_type=error_type,
        error_code=error_code,
        actual_value=actual_value,
        expected_value=expected_value,
        reason=f"{field_name}校验失败",
        source=(
            ValidationSource.COMPARISON
            if error_type is not ValidationErrorType.FORMAT_ERROR
            else ValidationSource.FORMAT_CHECK
        ),
    )


def _demo_result(source_file: Path) -> ValidationResult:
    return ValidationResult(
        file=source_file,
        summary=ValidationSummary(
            total_checks=10,
            pass_count=6,
            error_count=4,
            format_error_count=2,
            data_not_found_count=1,
            data_error_count=1,
            skipped_empty_row_count=3,
        ),
        issues=(
            _issue(
                source_file=source_file,
                sheet_name="设备档案",
                row_number=125,
                column_number=7,
                column_letter="G",
                field_name="设备类型编码",
                error_type=ValidationErrorType.DATA_ERROR,
                error_code="DATA_002",
                actual_value="D-CQJ",
                expected_value="D-CCQ",
            ),
            _issue(
                source_file=source_file,
                sheet_name="设备档案",
                row_number=128,
                column_number=3,
                column_letter="C",
                field_name="投用日期",
                error_type=ValidationErrorType.FORMAT_ERROR,
                error_code="FORMAT_001",
                actual_value="2026-13-45",
            ),
            _issue(
                source_file=source_file,
                sheet_name="技术参数",
                row_number=86,
                column_number=6,
                column_letter="F",
                field_name="数量",
                error_type=ValidationErrorType.DATA_NOT_FOUND,
                error_code="DATA_001",
                actual_value="D-XXX",
            ),
            _issue(
                source_file=source_file,
                sheet_name="技术参数",
                row_number=88,
                column_number=6,
                column_letter="F",
                field_name="数量",
                error_type=ValidationErrorType.FORMAT_ERROR,
                error_code="FORMAT_003",
                actual_value="1.5",
            ),
        ),
    )


def _read_sheets(path: Path):
    workbook = load_workbook(path, data_only=True)
    try:
        return workbook.sheetnames, workbook["核验汇总"], workbook["错误明细"]
    finally:
        workbook.close()


def test_export_creates_summary_and_detail_sheets(tmp_path: Path) -> None:
    source_file = tmp_path / "source.xlsx"
    build_workbook(source_file, {"设备档案": [["名称"]]})
    result = _demo_result(source_file)
    output_path = tmp_path / "报告" / "核验报告.xlsx"

    returned = export_validation_report(result, output_path)

    assert returned == output_path
    assert output_path.is_file()
    sheet_names, summary_sheet, detail_sheet = _read_sheets(output_path)
    assert sheet_names == ["核验汇总", "错误明细"]

    summary_rows = list(summary_sheet.iter_rows(values_only=True))
    summary_values = dict(summary_rows[1:])
    assert summary_values["核验文件"] == "source.xlsx"
    assert summary_values["总检查数"] == 10
    assert summary_values["通过数"] == 6
    assert summary_values["错误数"] == 4
    assert summary_values["格式错误"] == 2
    assert summary_values["数据不存在"] == 1
    assert summary_values["数据错误"] == 1
    assert summary_values["跳过空行"] == 3
    assert "4 条错误" in summary_values["结论"]

    detail_rows = list(detail_sheet.iter_rows(values_only=True))
    assert detail_rows[0] == (
        "序号",
        "文件",
        "Sheet",
        "行号",
        "列号",
        "列字母",
        "字段名",
        "单元格",
        "错误类型",
        "错误编码",
        "实际值",
        "期望值",
        "错误原因",
        "来源",
    )
    assert len(detail_rows) == 5


def test_detail_rows_keep_validation_locations_and_order(tmp_path: Path) -> None:
    source_file = tmp_path / "source.xlsx"
    result = _demo_result(source_file)
    output_path = tmp_path / "report.xlsx"

    export_validation_report(result, output_path)
    _, _, detail_sheet = _read_sheets(output_path)

    rows = list(detail_sheet.iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == [1, 2, 3, 4]
    assert [(row[1], row[2], row[3], row[4], row[5], row[6], row[7])
            for row in rows] == [
        ("source.xlsx", "设备档案", 125, 7, "G", "设备类型编码", "G125"),
        ("source.xlsx", "设备档案", 128, 3, "C", "投用日期", "C128"),
        ("source.xlsx", "技术参数", 86, 6, "F", "数量", "F86"),
        ("source.xlsx", "技术参数", 88, 6, "F", "数量", "F88"),
    ]
    assert [row[8:11] for row in rows] == [
        ("DATA_ERROR", "DATA_002", "D-CQJ"),
        ("FORMAT_ERROR", "FORMAT_001", "2026-13-45"),
        ("DATA_NOT_FOUND", "DATA_001", "D-XXX"),
        ("FORMAT_ERROR", "FORMAT_003", "1.5"),
    ]
    assert [row[11] for row in rows] == ["D-CCQ", None, None, None]
    assert [row[12] for row in rows] == [
        "设备类型编码校验失败",
        "投用日期校验失败",
        "数量校验失败",
        "数量校验失败",
    ]
    assert [row[13] for row in rows] == [
        "COMPARISON",
        "FORMAT_CHECK",
        "COMPARISON",
        "FORMAT_CHECK",
    ]


def test_all_passed_result_exports_summary_with_empty_detail(tmp_path: Path) -> None:
    source_file = tmp_path / "source.xlsx"
    result = ValidationResult(
        file=source_file,
        summary=ValidationSummary(
            total_checks=5,
            pass_count=5,
            error_count=0,
        ),
    )
    output_path = tmp_path / "report.xlsx"

    export_validation_report(result, output_path)
    _, summary_sheet, detail_sheet = _read_sheets(output_path)

    summary_values = dict(
        list(summary_sheet.iter_rows(values_only=True))[1:]
    )
    assert summary_values["结论"] == "全部通过，无错误。"
    assert summary_values["错误数"] == 0
    assert list(detail_sheet.iter_rows(values_only=True))[0][0] == "序号"
    assert detail_sheet.max_row == 1


def test_export_does_not_modify_source_file(tmp_path: Path) -> None:
    source_file = tmp_path / "source.xlsx"
    build_workbook(
        source_file,
        {"设备档案": [["名称"], ["数据行"]]},
    )
    before = source_file.read_bytes()
    result = _demo_result(source_file)
    output_path = tmp_path / "report.xlsx"

    export_validation_report(result, output_path)

    assert source_file.read_bytes() == before


def test_export_refuses_to_overwrite_source_file(tmp_path: Path) -> None:
    source_file = tmp_path / "source.xlsx"
    result = _demo_result(source_file)

    with pytest.raises(ValueError, match="不能覆盖源文件"):
        export_validation_report(result, source_file)


def test_report_output_keeps_source_values_as_text(tmp_path: Path) -> None:
    source_file = tmp_path / "source.xlsx"
    result = _demo_result(source_file)
    output_path = tmp_path / "report.xlsx"

    export_validation_report(result, output_path)
    _, _, detail_sheet = _read_sheets(output_path)

    assert detail_sheet["J2"].value == "DATA_002"
    assert detail_sheet["K2"].value == "D-CQJ"
    assert detail_sheet["L2"].value == "D-CCQ"
