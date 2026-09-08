"""P1-06 标准模板生成引擎测试。"""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_qc.cleaning import (
    CleaningWorkbookConfig,
    CleaningWorkbookResult,
    FieldCleaningRule,
    clean_workbook,
)
from excel_qc.equipment_type_mapping import (
    EquipmentTypeCodeConfig,
    EquipmentTypeRule,
    map_equipment_type_codes,
)
from excel_qc.field_mapping import (
    FieldMappingConfig,
    StandardFieldDefinition,
    WorkbookFieldMappingResult,
    map_workbook_fields,
)
from excel_qc.profiler import profile_workbook
from excel_qc.template_generator import (
    TemplateCellIssue,
    TemplateColumnConfig,
    TemplateConfig,
    TemplateConfigError,
    TemplateDataError,
    TemplateGenerationError,
    TemplateIssueSourceStage,
    TemplateRow,
    TemplateWorkbookResult,
    build_standard_template,
    export_standard_template,
    load_template_config,
)
from tests.qc.helpers import build_workbook


_BASIC_FIELDS = (
    StandardFieldDefinition(name="设备名称"),
    StandardFieldDefinition(name="设备类型名称", aliases=("设备类型",)),
    StandardFieldDefinition(name="备注"),
)

_DEFAULT_EQUIPMENT_RULES = (
    EquipmentTypeRule(
        aliases=("除尘器",),
        standard_type_name="除尘器",
        system_code="D-CCQ",
    ),
    EquipmentTypeRule(
        aliases=("空压机",),
        standard_type_name="空压机",
        system_code="D-KYJ",
    ),
)


def _equipment_config(
    *,
    target: str = "设备类型名称",
    rules: tuple[EquipmentTypeRule, ...] = _DEFAULT_EQUIPMENT_RULES,
    code_pattern: str | None = None,
) -> EquipmentTypeCodeConfig:
    return EquipmentTypeCodeConfig(
        target_standard_field=target,
        rules=rules,
        code_pattern=code_pattern,
    )


def _run_chain(
    tmp_path: Path,
    filename: str,
    sheets: dict[str, list[list[object]]],
    *,
    standard_fields: tuple[StandardFieldDefinition, ...] = _BASIC_FIELDS,
    cleaning_rules: tuple[FieldCleaningRule, ...] = (),
    equipment_config: EquipmentTypeCodeConfig | None = None,
):
    path = build_workbook(tmp_path / filename, sheets)
    profile = profile_workbook(path)
    mapping = map_workbook_fields(
        profile, FieldMappingConfig(standard_fields)
    )
    cleaning = clean_workbook(
        profile,
        mapping,
        CleaningWorkbookConfig(cleaning_rules),
    )
    config = (
        equipment_config
        if equipment_config is not None
        else _equipment_config()
    )
    equipment = map_equipment_type_codes(cleaning, config)
    return path, profile, mapping, cleaning, equipment


def _config(
    *columns: TemplateColumnConfig,
    sheets: tuple[str, ...] | None = None,
) -> TemplateConfig:
    return TemplateConfig(columns=tuple(columns), sheets=sheets)


def _name_column(header: str = "设备名称") -> TemplateColumnConfig:
    return TemplateColumnConfig(standard_field="设备名称", header=header)


def _code_column(header: str = "设备类型编码") -> TemplateColumnConfig:
    return TemplateColumnConfig(
        standard_field="设备类型编码",
        header=header,
        code_source_standard_field="设备类型名称",
    )


def _write_json(tmp_path: Path, payload: object) -> Path:
    config_path = tmp_path / "template.json"
    config_path.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )
    return config_path


def test_load_template_config_valid_json(tmp_path: Path) -> None:
    config_path = _write_json(
        tmp_path,
        {
            "columns": [
                {"standard_field": "设备名称", "header": "设备名称"},
                {
                    "standard_field": "设备类型编码",
                    "header": "设备类型编码",
                    "code_source_standard_field": "设备类型名称",
                },
            ]
        },
    )

    config = load_template_config(config_path)

    assert len(config.columns) == 2
    assert config.columns[0].standard_field == "设备名称"
    assert config.columns[0].output_header == "设备名称"
    assert config.columns[1].is_code_column
    assert config.columns[1].code_source_standard_field == "设备类型名称"
    assert config.sheets is None


def test_load_template_config_header_defaults_to_standard_field(
    tmp_path: Path,
) -> None:
    config_path = _write_json(
        tmp_path,
        {"columns": [{"standard_field": "设备名称"}]},
    )

    config = load_template_config(config_path)

    assert config.columns[0].output_header == "设备名称"


def test_load_template_config_sheets(tmp_path: Path) -> None:
    config_path = _write_json(
        tmp_path,
        {
            "sheets": ["设备档案"],
            "columns": [{"standard_field": "设备名称", "header": "设备名称"}],
        },
    )

    config = load_template_config(config_path)

    assert config.sheets == ("设备档案",)


def test_load_template_config_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(tmp_path / "missing.json")

    assert "不存在" in str(exc_info.value)


@pytest.mark.parametrize(
    "payload",
    [
        "{bad json",
        "[]",
        '{"columns": 1}',
    ],
)
def test_load_template_config_invalid_json_structure_raises(
    tmp_path: Path,
    payload: str,
) -> None:
    config_path = tmp_path / "bad.json"
    config_path.write_text(payload, encoding="utf-8")

    with pytest.raises(TemplateConfigError):
        load_template_config(config_path)


def test_load_template_config_missing_columns_raises(tmp_path: Path) -> None:
    config_path = _write_json(tmp_path, {"sheets": ["设备档案"]})

    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(config_path)

    assert "columns" in str(exc_info.value)


def test_load_template_config_empty_columns_raises(tmp_path: Path) -> None:
    config_path = _write_json(tmp_path, {"columns": []})

    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(config_path)

    assert "columns" in str(exc_info.value)


@pytest.mark.parametrize(
    "payload",
    [
        {"columns": [{"standard_field": "", "header": "设备名称"}]},
        {"columns": [{"standard_field": 5, "header": "设备名称"}]},
    ],
)
def test_load_template_config_invalid_standard_field_raises(
    tmp_path: Path,
    payload: dict,
) -> None:
    config_path = _write_json(tmp_path, payload)

    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(config_path)

    assert "standard_field" in str(exc_info.value)


@pytest.mark.parametrize(
    "payload",
    [
        {"columns": [{"standard_field": "设备名称", "header": ""}]},
        {"columns": [{"standard_field": "设备名称", "header": 3}]},
    ],
)
def test_load_template_config_invalid_header_raises(
    tmp_path: Path,
    payload: dict,
) -> None:
    config_path = _write_json(tmp_path, payload)

    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(config_path)

    assert "header" in str(exc_info.value)


@pytest.mark.parametrize(
    "payload",
    [
        {
            "columns": [
                {"standard_field": "设备名称", "header": "设备名称"},
                {"standard_field": "设备名称", "header": "名称"},
            ]
        },
        {
            "columns": [
                {"standard_field": "设备名称", "header": "名称"},
                {"standard_field": "备注", "header": "名称"},
            ]
        },
    ],
)
def test_load_template_config_duplicate_column_raises(
    tmp_path: Path,
    payload: dict,
) -> None:
    config_path = _write_json(tmp_path, payload)

    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(config_path)

    assert "重复" in str(exc_info.value)


@pytest.mark.parametrize(
    "code_source",
    ["", 5],
)
def test_load_template_config_incomplete_code_source_raises(
    tmp_path: Path,
    code_source: object,
) -> None:
    config_path = _write_json(
        tmp_path,
        {
            "columns": [
                {
                    "standard_field": "设备类型编码",
                    "header": "设备类型编码",
                    "code_source_standard_field": code_source,
                }
            ]
        },
    )

    with pytest.raises(TemplateConfigError) as exc_info:
        load_template_config(config_path)

    assert "code_source_standard_field" in str(exc_info.value)


@pytest.mark.parametrize(
    "sheets",
    [[], ["设备档案", "设备档案"]],
)
def test_load_template_config_invalid_sheets_raises(
    tmp_path: Path,
    sheets: list[str],
) -> None:
    config_path = _write_json(
        tmp_path,
        {
            "sheets": sheets,
            "columns": [{"standard_field": "设备名称", "header": "设备名称"}],
        },
    )

    with pytest.raises(TemplateConfigError):
        load_template_config(config_path)


def test_single_sheet_assembly_column_order_and_values(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "single.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型", "备注"],
                [" 泵A ", "除尘器", "第一台"],
                ["风机", "空压机", "第二台"],
            ]
        },
        cleaning_rules=(FieldCleaningRule(standard_field="设备名称"),),
    )
    template_config = _config(
        TemplateColumnConfig(standard_field="备注", header="备注信息"),
        _name_column(),
        _code_column(),
    )

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    sheet_result = result.sheet_results[0]
    assert len(sheet_result.rows) == 2
    assert sheet_result.sheet_name == "设备档案"
    assert [column.output_header for column in sheet_result.columns] == [
        "备注信息",
        "设备名称",
        "设备类型编码",
    ]
    assert sheet_result.rows[0].source_row_number == 2
    assert sheet_result.rows[0].values == ("第一台", "泵A", "D-CCQ")
    assert sheet_result.rows[1].values == ("第二台", "风机", "D-KYJ")
    assert result.summary.sheet_count == 1
    assert result.summary.row_count == 2
    assert result.summary.cell_count == 6
    assert result.summary.issue_count == 0
    assert result.issues == ()


def test_multi_sheet_assembly_preserves_profile_order(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "multi.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ],
            "补充资料": [
                ["设备名称", "设备类型"],
                ["风机", "空压机"],
            ],
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    assert [item.sheet_name for item in result.sheet_results] == [
        "设备档案",
        "补充资料",
    ]
    assert result.sheet_results[0].rows[0].source_row_number == 2
    assert result.sheet_results[1].rows[0].source_row_number == 2
    assert result.sheet_results[0].rows[0].values[1] == "D-CCQ"
    assert result.sheet_results[1].rows[0].values[1] == "D-KYJ"


def test_inner_blank_rows_are_preserved(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "blank.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
                [None, None],
                ["风机", "空压机"],
            ]
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    rows = result.sheet_results[0].rows
    assert [row.source_row_number for row in rows] == [2, 3, 4]
    assert rows[0].values == ("泵A", "D-CCQ")
    assert rows[1].values == (None, None)
    assert rows[2].values == ("风机", "D-KYJ")
    assert rows[1].trace[0].source_column_letter == "A"
    assert rows[1].trace[0].original_value is None


def test_missing_standard_field_creates_blank_column_with_single_issue(
    tmp_path: Path,
) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "missing-field.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
                ["风机", "空压机"],
            ]
        },
    )
    template_config = _config(
        _name_column(),
        TemplateColumnConfig(standard_field="设备编码", header="设备编码"),
    )

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    sheet_result = result.sheet_results[0]
    assert all(row.values[1] is None for row in sheet_result.rows)
    missing_issues = [
        issue
        for issue in sheet_result.issues
        if issue.issue_code == "FIELD_MAPPING_MISSING"
    ]
    assert len(missing_issues) == 1
    assert missing_issues[0].sheet == "设备档案"
    assert missing_issues[0].source_row_number is None
    assert missing_issues[0].standard_field == "设备编码"
    assert missing_issues[0].template_column == "设备编码"
    assert missing_issues[0].source_stage is TemplateIssueSourceStage.FIELD_MAPPING
    assert "不做任何猜测" in missing_issues[0].reason


def test_field_mapping_conflict_column_is_blank_with_issue(
    tmp_path: Path,
) -> None:
    fields = (
        StandardFieldDefinition(
            name="设备X", aliases=("设备名称", "名称")
        ),
    )
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "conflict-field.xlsx",
        {
            "设备档案": [
                ["设备名称", "名称"],
                ["a", "b"],
            ]
        },
        standard_fields=fields,
        equipment_config=_equipment_config(target="设备X", rules=()),
    )
    template_config = _config(
        TemplateColumnConfig(standard_field="设备X", header="设备X")
    )

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    sheet_result = result.sheet_results[0]
    assert sheet_result.rows[0].values == (None,)
    conflict_issues = [
        issue
        for issue in sheet_result.issues
        if issue.issue_code == "FIELD_MAPPING_CONFLICT"
    ]
    assert len(conflict_issues) == 1
    assert "需人工确认" in conflict_issues[0].reason


def test_code_column_writes_match_system_code(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "code-match.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    row = result.sheet_results[0].rows[0]
    assert row.values[1] == "D-CCQ"
    assert row.trace[1].system_code == "D-CCQ"
    assert result.summary.issue_count == 0


@pytest.mark.parametrize(
    ("value", "issue_code", "reason_text"),
    [
        ("未知类型", "EQUIPMENT_UNMATCHED", "未命中"),
        (None, "EQUIPMENT_UNMATCHED", "空值"),
    ],
)
def test_code_column_unmatched_is_blank_with_issue(
    tmp_path: Path,
    value: object,
    issue_code: str,
    reason_text: str,
) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "code-unmatched.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", value],
            ]
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    row = result.sheet_results[0].rows[0]
    assert row.values[1] is None
    assert row.trace[1].system_code is None
    code_issues = [
        issue
        for issue in result.sheet_results[0].issues
        if issue.template_column == "设备类型编码"
    ]
    assert len(code_issues) == 1
    assert code_issues[0].issue_code == issue_code
    assert reason_text in code_issues[0].reason
    assert code_issues[0].source_row_number == 2


def test_code_column_conflict_is_blank_with_issue(tmp_path: Path) -> None:
    config = _equipment_config(
        rules=(
            EquipmentTypeRule(
                aliases=("COMMON",),
                standard_type_name="类型A",
                system_code="A-001",
            ),
            EquipmentTypeRule(
                aliases=("COMMON",),
                standard_type_name="类型B",
                system_code="B-001",
            ),
        )
    )
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "code-conflict.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "COMMON"],
            ]
        },
        equipment_config=config,
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    row = result.sheet_results[0].rows[0]
    assert row.values[1] is None
    issue = result.sheet_results[0].issues[0]
    assert issue.issue_code == "EQUIPMENT_CONFLICT"
    assert "命中多个规则" in issue.reason
    assert issue.source_stage is TemplateIssueSourceStage.EQUIPMENT_TYPE_MAPPING


def test_code_column_invalid_is_blank_with_issue(tmp_path: Path) -> None:
    config = _equipment_config(
        code_pattern=r"^D-[A-Z0-9]{3}$",
        rules=(
            EquipmentTypeRule(
                aliases=("除尘器",),
                standard_type_name="除尘器",
                system_code="BAD-CODE",
            ),
        ),
    )
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "code-invalid.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
        equipment_config=config,
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    row = result.sheet_results[0].rows[0]
    assert row.values[1] is None
    issue = result.sheet_results[0].issues[0]
    assert issue.issue_code == "EQUIPMENT_INVALID"
    assert "code_pattern" in issue.reason
    assert row.trace[1].system_code == "BAD-CODE"
    assert row.trace[1].cleaned_value == "除尘器"


def test_code_column_skipped_when_target_field_missing(tmp_path: Path) -> None:
    fields = (
        StandardFieldDefinition(name="设备名称"),
        StandardFieldDefinition(name="备注"),
    )
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "code-skipped.xlsx",
        {
            "设备档案": [
                ["设备名称", "备注"],
                ["泵A", "备注值"],
            ]
        },
        standard_fields=fields,
        equipment_config=_equipment_config(target="设备类型名称", rules=()),
    )
    template_config = _config(
        _name_column(),
        _code_column(),
        TemplateColumnConfig(standard_field="备注", header="备注"),
    )

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    row = result.sheet_results[0].rows[0]
    assert row.values[1] is None
    skipped_issues = [
        issue
        for issue in result.sheet_results[0].issues
        if issue.issue_code == "EQUIPMENT_SKIPPED"
    ]
    assert len(skipped_issues) == 1
    assert "未找到目标标准字段" in skipped_issues[0].reason


def test_normal_field_trace_full_source_chain(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "trace-normal.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["  泵A  ", "除尘器"],
            ]
        },
        cleaning_rules=(FieldCleaningRule(standard_field="设备名称"),),
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    trace = result.sheet_results[0].rows[0].trace[0]
    assert trace.sheet == "设备档案"
    assert trace.source_row_number == 2
    assert trace.source_column_number == 1
    assert trace.source_column_letter == "A"
    assert trace.source_field_name == "设备名称"
    assert trace.standard_field == "设备名称"
    assert trace.template_column == "设备名称"
    assert trace.original_value == "  泵A  "
    assert trace.cleaned_value == "泵A"
    assert trace.system_code is None


def test_code_field_trace_full_source_chain(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "trace-code.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    trace = result.sheet_results[0].rows[0].trace[1]
    assert trace.sheet == "设备档案"
    assert trace.source_row_number == 2
    assert trace.source_column_number == 2
    assert trace.source_column_letter == "B"
    assert trace.source_field_name == "设备类型"
    assert trace.standard_field == "设备类型名称"
    assert trace.template_column == "设备类型编码"
    assert trace.original_value == "除尘器"
    assert trace.cleaned_value == "除尘器"
    assert trace.system_code == "D-CCQ"


def test_issue_reason_is_preserved_from_equipment_audit(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "issue-reason.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "未知类型"],
            ]
        },
    )
    template_config = _config(_code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    audit = equipment.sheet_results[0].audits[0]
    issue = result.sheet_results[0].issues[0]
    assert issue.reason == audit.reason
    assert issue.source_row_number == audit.row_number
    assert issue.standard_field == "设备类型名称"


def test_sheet_name_mismatch_raises_template_data_error(
    tmp_path: Path,
) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "sheet-mismatch.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ],
            "补充资料": [
                ["设备名称", "设备类型"],
                ["风机", "空压机"],
            ],
        },
    )
    reversed_mapping = WorkbookFieldMappingResult(
        file_path=mapping.file_path,
        sheet_results=tuple(reversed(mapping.sheet_results)),
        summary=mapping.summary,
    )

    with pytest.raises(TemplateDataError) as exc_info:
        build_standard_template(
            profile,
            reversed_mapping,
            cleaning,
            equipment,
            _config(_name_column()),
        )

    assert "Sheet" in str(exc_info.value)


def test_row_range_mismatch_raises_template_data_error(
    tmp_path: Path,
) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "row-mismatch.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
                ["风机", "空压机"],
            ]
        },
    )
    first_sheet = cleaning.sheet_results[0]
    modified_sheet = replace(first_sheet, rows=first_sheet.rows[:-1])
    modified_cleaning = CleaningWorkbookResult(
        source_file=cleaning.source_file,
        sheet_results=(modified_sheet, *cleaning.sheet_results[1:]),
        summary=cleaning.summary,
    )

    with pytest.raises(TemplateDataError) as exc_info:
        build_standard_template(
            profile,
            mapping,
            modified_cleaning,
            equipment,
            _config(_name_column()),
        )

    assert "行覆盖范围不一致" in str(exc_info.value)


def test_sheet_filter_in_config_selects_subset(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "select.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ],
            "补充资料": [
                ["设备名称", "设备类型"],
                ["风机", "空压机"],
            ],
        },
    )
    template_config = _config(
        _name_column(), sheets=("补充资料",)
    )

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    assert [item.sheet_name for item in result.sheet_results] == ["补充资料"]


def test_build_does_not_modify_input_objects(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "pure.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    profile_before = profile
    mapping_before = mapping
    cleaning_before = cleaning
    equipment_before = equipment

    build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column(), _code_column()),
    )

    assert profile == profile_before
    assert mapping == mapping_before
    assert cleaning == cleaning_before
    assert equipment == equipment_before


def test_build_creates_no_files(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "no-write.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    before_files = {path.name for path in tmp_path.iterdir()}
    source_bytes = (tmp_path / "no-write.xlsx").read_bytes()

    build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column(), _code_column()),
    )

    assert {path.name for path in tmp_path.iterdir()} == before_files
    assert (tmp_path / "no-write.xlsx").read_bytes() == source_bytes


def test_export_writes_xlsx_with_headers_and_values(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "export.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
                ["风机", "空压机"],
            ]
        },
        cleaning_rules=(FieldCleaningRule(standard_field="设备名称"),),
    )
    result = build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column(), _code_column()),
    )
    output_path = tmp_path / "out" / "standard.xlsx"

    returned = export_standard_template(result, output_path)

    assert returned == output_path
    assert output_path.exists()
    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == ["设备档案"]
    sheet = workbook["设备档案"]
    assert sheet.max_row == 3
    assert sheet.freeze_panes == "A2"
    assert list(sheet.iter_rows(values_only=True)) == [
        ("设备名称", "设备类型编码"),
        ("泵A", "D-CCQ"),
        ("风机", "D-KYJ"),
    ]
    workbook.close()


def test_export_multiple_sheets_preserves_order_and_content(
    tmp_path: Path,
) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "export-multi.xlsx",
        {
            "甲表": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ],
            "乙表": [
                ["设备名称", "设备类型"],
                ["风机", "空压机"],
            ],
        },
    )
    result = build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column(), _code_column()),
    )
    output_path = tmp_path / "multi-out.xlsx"

    export_standard_template(result, output_path)

    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == ["甲表", "乙表"]
    assert workbook["甲表"]["A2"].value == "泵A"
    assert workbook["乙表"]["A2"].value == "风机"
    workbook.close()


def test_export_refuses_to_overwrite_source_file(tmp_path: Path) -> None:
    source_path, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "same-path.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    result = build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column()),
    )
    source_bytes = source_path.read_bytes()

    with pytest.raises(TemplateGenerationError) as exc_info:
        export_standard_template(result, source_path)

    assert "不能覆盖源文件" in str(exc_info.value)
    assert source_path.read_bytes() == source_bytes


def test_export_creates_parent_directory(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "nested.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    result = build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column()),
    )
    output_path = tmp_path / "a" / "b" / "c" / "模板.xlsx"

    export_standard_template(result, output_path)

    assert output_path.exists()
    assert output_path.parent.is_dir()


def test_export_accepts_result_and_returns_path_object(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "export-return.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    result = build_standard_template(
        profile,
        mapping,
        cleaning,
        equipment,
        _config(_name_column()),
    )
    output_path = Path(tmp_path / "returned.xlsx")

    returned = export_standard_template(result, output_path)

    assert isinstance(returned, Path)
    assert returned == output_path


def test_summary_counts_rows_cells_and_issues(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "summary.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
                ["风机", "未知类型"],
            ]
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    assert result.summary.sheet_count == 1
    assert result.summary.row_count == 2
    assert result.summary.cell_count == 4
    assert result.summary.issue_count == 1
    assert isinstance(result.issues[0], TemplateCellIssue)


def test_result_model_contains_rows_with_ordered_values(tmp_path: Path) -> None:
    _, profile, mapping, cleaning, equipment = _run_chain(
        tmp_path,
        "models.xlsx",
        {
            "设备档案": [
                ["设备名称", "设备类型"],
                ["泵A", "除尘器"],
            ]
        },
    )
    template_config = _config(_name_column(), _code_column())

    result = build_standard_template(
        profile, mapping, cleaning, equipment, template_config
    )

    assert isinstance(result, TemplateWorkbookResult)
    row = result.sheet_results[0].rows[0]
    assert isinstance(row, TemplateRow)
    assert row.source_row_number == 2
    assert len(row.values) == len(row.trace) == 2
