"""P0 阶段测试专用：动态生成 xlsx 样例。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook


def build_workbook(
    path: Path,
    sheets: dict[str, list[list[Any]]],
) -> Path:
    """按 {Sheet 名: [行, ...]} 生成 .xlsx 文件。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for sheet_name, rows in sheets.items():
        worksheet = book.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(list(row))
    book.save(path)
    return path


def build_typed_workbook(path: Path) -> Path:
    """生成含多种基础数据类型的样例。"""
    return build_workbook(
        path,
        {
            "设备档案": [
                ["设备编码", "设备名称", "数量", "单价", "投用日期", "在用", "备注"],
                ["EQ-001", "空压机", 2, 12.5, datetime(2024, 1, 5, 9, 30), True, None],
                ["EQ-002", "冷水机", 3, 8.0, datetime(2024, 2, 10), False, None],
            ]
        },
    )


def build_workbook_with_metadata(
    path: Path,
    sheets: dict[str, list[list[Any]]],
    *,
    merged_ranges: dict[str, list[str]] | None = None,
    hidden_rows: dict[str, list[int]] | None = None,
    hidden_columns: dict[str, list[str]] | None = None,
) -> Path:
    """生成可设置合并单元格与隐藏行列的 .xlsx 样例。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for sheet_name, rows in sheets.items():
        worksheet = book.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(list(row))
        for cell_range in (merged_ranges or {}).get(sheet_name, []):
            worksheet.merge_cells(cell_range)
        for row_number in (hidden_rows or {}).get(sheet_name, []):
            worksheet.row_dimensions[row_number].hidden = True
        for column_letter in (hidden_columns or {}).get(sheet_name, []):
            worksheet.column_dimensions[column_letter].hidden = True
    book.save(path)
    return path
