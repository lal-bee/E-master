"""Excel 文件导入（P0-01）。

当前阶段只支持 .xlsx。读取过程使用 openpyxl 只读模式，
不会修改用户的原始文件。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook as _openpyxl_load_workbook

from excel_qc.errors import (
    SourceFileNotFoundError,
    UnsupportedFileTypeError,
    WorkbookReadError,
)


@dataclass(frozen=True)
class LoadedSheet:
    """一个 Sheet 的原始读取结果。

    rows 保留 Excel 物理行顺序：索引 0 对应 Excel 第 1 行。
    include_sheet_metadata=True 时额外记录 Sheet 索引、物理总行列数、
    合并范围与隐藏行/列信息；默认读取路径不填充这些字段。
    """

    sheet_name: str
    rows: tuple[tuple[Any, ...], ...]
    sheet_index: int = 0
    total_row_count: int | None = None
    total_column_count: int | None = None
    merged_cell_ranges: tuple[str, ...] = ()
    hidden_row_numbers: tuple[int, ...] = ()
    hidden_column_letters: tuple[str, ...] = ()


@dataclass(frozen=True)
class LoadedWorkbook:
    """一个 Excel 文件的原始读取结果。"""

    path: Path
    file_type: str
    sheets: tuple[LoadedSheet, ...]

    @property
    def sheet_count(self) -> int:
        return len(self.sheets)

    @property
    def sheet_names(self) -> tuple[str, ...]:
        return tuple(sheet.sheet_name for sheet in self.sheets)


def load_workbook(
    path: str | Path,
    *,
    include_sheet_metadata: bool = False,
) -> LoadedWorkbook:
    """只读导入 .xlsx 文件并返回全部 Sheet 的原始内容。

    include_sheet_metadata=False（默认）保持 V1.0 只读模式行为；
    include_sheet_metadata=True 时以常规模式读取以获取合并单元格与
    隐藏行列元数据，仅用于结构探查，不修改原文件。
    """
    source = Path(path)
    if not source.exists() or not source.is_file():
        raise SourceFileNotFoundError(f"源文件不存在: {source}")
    if source.suffix.lower() != ".xlsx":
        raise UnsupportedFileTypeError(
            f"不支持的文件类型 '{source.suffix}'，P0 阶段仅支持 .xlsx"
        )

    workbook = None
    try:
        workbook = _openpyxl_load_workbook(
            source,
            read_only=not include_sheet_metadata,
            data_only=True,
        )
        sheets = tuple(
            LoadedSheet(
                sheet_name=worksheet.title,
                rows=tuple(tuple(row) for row in worksheet.iter_rows(values_only=True)),
                sheet_index=(index if include_sheet_metadata else 0),
                total_row_count=(
                    worksheet.max_row if include_sheet_metadata else None
                ),
                total_column_count=(
                    worksheet.max_column if include_sheet_metadata else None
                ),
                merged_cell_ranges=(
                    tuple(
                        str(cell_range)
                        for cell_range in worksheet.merged_cells.ranges
                    )
                    if include_sheet_metadata
                    else ()
                ),
                hidden_row_numbers=(
                    _hidden_row_numbers(worksheet)
                    if include_sheet_metadata
                    else ()
                ),
                hidden_column_letters=(
                    _hidden_column_letters(worksheet)
                    if include_sheet_metadata
                    else ()
                ),
            )
            for index, worksheet in enumerate(workbook.worksheets)
        )
    except Exception as exc:  # noqa: BLE001 - 文件层异常统一包装，避免底层异常泄漏
        raise WorkbookReadError(f"无法读取 Excel 文件 {source}：{exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()

    return LoadedWorkbook(
        path=source,
        file_type=source.suffix.lower().lstrip("."),
        sheets=sheets,
    )


def _hidden_row_numbers(worksheet) -> tuple[int, ...]:
    numbers = [
        int(row_number)
        for row_number, dimension in worksheet.row_dimensions.items()
        if dimension.hidden
    ]
    return tuple(sorted(numbers))


def _hidden_column_letters(worksheet) -> tuple[str, ...]:
    letters = [
        str(column_letter)
        for column_letter, dimension in worksheet.column_dimensions.items()
        if dimension.hidden
    ]
    return tuple(sorted(letters))
