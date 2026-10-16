"""固定字段规则；映射来自既有 Excel，不推断业务编码。"""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel

DATE_FIELDS = ("出厂日期", "启用日期", "设备制造日期")
NAME_FIELD = "设备小类名称"
CODE_FIELD = "设备小类编码"
DEVICE_FIELD = "设备代码"
TARGET_FIELDS = (NAME_FIELD, CODE_FIELD, *DATE_FIELDS, DEVICE_FIELD)
DEFAULT_MAPPING = Path(__file__).resolve().parents[1] / "docs" / "设备小类编码字段映射配置文件.xlsx"


class MappingError(ValueError):
    pass


class YearOnlyError(ValueError):
    pass


def load_mapping(path: Path = DEFAULT_MAPPING) -> dict[str, str]:
    """读取严格的一一映射；配置问题完整汇总，不采用最后一条覆盖。"""
    problems, mapping, codes = [], {}, {}
    try:
        book = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise MappingError(f"无法读取设备小类映射文件：{path.name}") from exc
    try:
        found = False
        for sheet in book:
            rows = iter(sheet.iter_rows())
            header = next(rows, ())
            fields = [cell.value for cell in header]
            if NAME_FIELD not in fields or CODE_FIELD not in fields:
                continue
            found = True
            if fields.count(NAME_FIELD) != 1 or fields.count(CODE_FIELD) != 1:
                problems.append(f"{sheet.title}：映射表头重复")
                continue
            ni, ci = fields.index(NAME_FIELD), fields.index(CODE_FIELD)
            for number, cells in enumerate(rows, 2):
                name, code = cells[ni], cells[ci]
                if name.value is None and code.value is None:
                    continue
                if name.data_type == "f" or code.data_type == "f":
                    problems.append(f"{sheet.title} 第 {number} 行：映射不能使用公式")
                    continue
                n = "" if name.value is None else str(name.value)
                c = "" if code.value is None else str(code.value)
                if not n.strip() or not c.strip():
                    problems.append(f"{sheet.title} 第 {number} 行：名称或编码为空")
                elif n in mapping and mapping[n] != c:
                    problems.append(f"{sheet.title} 第 {number} 行：{n} 对应多个编码")
                elif c in codes and codes[c] != n:
                    problems.append(f"{sheet.title} 第 {number} 行：{c} 对应多个名称")
                else:
                    mapping[n], codes[c] = c, n
        if not found or not mapping:
            problems.append("映射文件缺少有效的设备小类名称／设备小类编码表")
    finally:
        book.close()
    if problems:
        raise MappingError("；".join(problems))
    return mapping


def normalize_date(value, *, numeric=False, epoch=WINDOWS_EPOCH) -> str | None:
    """三列统一入口：先判断年/月/日，最后才考虑真实数值序列。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return None
    text = str(value).strip()
    if re.fullmatch(r"\d{4}(?:\.0+)?", text):
        raise YearOnlyError(text)
    match = re.fullmatch(r"(\d{4})([-/.])(\d{1,2})(?:\2(\d{1,2}))?", text)
    if match:
        try:
            return date(int(match[1]), int(match[3]), int(match[4] or 1)).isoformat()
        except ValueError:
            return None
    # 时间字符串与 OOXML 原生日期类型。
    if not numeric:
        try:
            return datetime.fromisoformat(text).date().isoformat()
        except ValueError:
            return None
    try:
        serial = Decimal(text)
        if not serial.is_finite():
            return None
        if serial == serial.to_integral_value() and 1000 <= serial <= 9999:
            raise YearOnlyError(text)
        # 科学计数法数值也不能让四位年份绕过检查。
        plain = format(serial, "f")
        if re.fullmatch(r"\d{4}\.\d{1,2}", plain):
            return normalize_date(plain)
        if epoch == MAC_EPOCH and 0 <= serial < 1:
            return epoch.date().isoformat()
        if serial < 1:
            return None
        parsed = from_excel(float(serial), epoch)
        return parsed.date().isoformat() if isinstance(parsed, datetime) else None
    except (InvalidOperation, ValueError, OverflowError):
        return None


def normalize_device_code(value, *, numeric=False, number_format="") -> str | None:
    if value is None:
        return None
    text = str(value)
    if text.strip().lower() in {"", "null", "/"}:
        return None
    if numeric:
        try:
            number = Decimal(text)
            if not number.is_finite():
                return None
            text = format(number, "f")
            if number == number.to_integral_value():
                text = text.split(".")[0]
                if re.fullmatch(r"0+", number_format):
                    text = text.zfill(len(number_format))
        except InvalidOperation:
            return None
    # 已处理的合法结果再次执行时保持稳定。
    try:
        parsed = json.loads(text)
        if isinstance(parsed, dict) and set(parsed) == {"注册代码"} and isinstance(parsed["注册代码"], str):
            if parsed["注册代码"].strip().lower() in {"", "null", "/"}:
                return None
            text = parsed["注册代码"]
    except (ValueError, TypeError):
        pass
    return json.dumps({"注册代码": text}, ensure_ascii=False, separators=(",", ":"))
