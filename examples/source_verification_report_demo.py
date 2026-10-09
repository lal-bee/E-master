"""P1-10 演示：临时双 Excel → P1-09 核验 → P1-10 报告；非真实业务模板。"""

from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from examples.source_verification_demo import demo_config  # noqa: E402
from excel_qc import verify_source_against_system_export  # noqa: E402
from ui import export_source_verification_report  # noqa: E402


def _write(path: Path, header: tuple[str, str], rows: tuple[tuple[object, object], ...]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "演示数据"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def main() -> None:
    with TemporaryDirectory(prefix="p1-10-demo-") as directory:
        root = Path(directory)
        source = root / "original.xlsx"
        system = root / "system.xlsx"
        report = root / "verification-report.xlsx"
        _write(source, ("源编号", "源名称"), (
            ("MATCH", "同值"), ("CHANGED", "旧值"), ("MISSING", "仅源"),
            ("DUP", "甲"), ("DUP", "乙"), (None, "无主键"),
        ))
        _write(system, ("系统编号", "系统名称"), (
            ("MATCH", "同值"), ("CHANGED", "新值"), ("EXTRA", "仅系统"), ("DUP", "甲"),
        ))
        verification = verify_source_against_system_export(source, system, demo_config())
        export_source_verification_report(verification, report)
        workbook = load_workbook(report, read_only=True)
        try:
            print("Sheets:", ", ".join(workbook.sheetnames))
            for name in workbook.sheetnames:
                print(name, "rows:", workbook[name].max_row)
            print("verification_complete:", verification.summary.verification_complete)
            print("data_consistent:", verification.summary.data_consistent)
        finally:
            workbook.close()


if __name__ == "__main__":
    main()
